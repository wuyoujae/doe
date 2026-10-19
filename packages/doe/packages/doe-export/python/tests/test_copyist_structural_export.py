from __future__ import annotations

import copy
import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from modular_patch import patch

from pptx import Presentation as PptxPresentation


RUNTIME_ROOT = Path(__file__).resolve().parents[1]
if str(RUNTIME_ROOT) not in sys.path:
    sys.path.insert(0, str(RUNTIME_ROOT))

import pptx_pipeline  # noqa: E402
from pptx_pipeline import export_ir_to_pptx, import_pptx_to_canonical  # noqa: E402
from pptx2html.server.pptx_writer import write_pptx_from_ir  # noqa: E402


def _shape(shape_id: int, name: str, left: float) -> dict:
    return {
        "type": "auto_shape",
        "shape_id": shape_id,
        "name": name,
        "geometry": "rect",
        "bounds": {"left": left, "top": 40, "width": 240, "height": 120},
        "fill": {"kind": "solid", "color": "4472C4"},
        "border": {"width": 1, "color": "203040"},
    }


def _slide(index: int, shape: dict) -> dict:
    return {
        "index": index,
        "name": f"Slide {index + 1}",
        "background": {"kind": "solid", "color": "FFFFFF"},
        "shapes": [shape],
    }


class CopyistStructuralExportTests(unittest.TestCase):
    def test_inserts_authored_slide_without_losing_imported_slide_order(self) -> None:
        original = {
            "width": 960,
            "height": 540,
            "slides": [
                _slide(0, _shape(2, "Imported First", 40)),
                _slide(1, _shape(3, "Imported Second", 320)),
            ],
        }
        desired = copy.deepcopy(original)
        desired["slides"] = [
            {**copy.deepcopy(original["slides"][0]), "_yrkie_source_slide_index": 0},
            _slide(1, _shape(10, "Inserted Authored", 180)),
            {**copy.deepcopy(original["slides"][1]), "_yrkie_source_slide_index": 1},
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "source.pptx"
            output_path = Path(temp_dir) / "output.pptx"
            write_pptx_from_ir(original, source_path)
            write_pptx_from_ir(
                desired,
                output_path,
                source_pptx=source_path.read_bytes(),
                original_ir=original,
            )

            presentation = PptxPresentation(output_path)
            self.assertEqual(len(presentation.slides), 3)
            left_positions = [
                round(slide.shapes[0].left / 12700)
                for slide in presentation.slides
            ]
            self.assertEqual(left_positions, [40, 180, 320])

    def test_export_normalization_keeps_source_provenance_and_scales_authored_slide(
        self,
    ) -> None:
        original = {
            "width": 720,
            "height": 405,
            "slides": [_slide(0, _shape(2, "Imported", 40))],
        }
        authored_shape = _shape(10, "Authored", 0)
        authored_shape["bounds"] = {
            "left": 0,
            "top": 0,
            "width": 960,
            "height": 540,
        }
        desired = {
            "width": 960,
            "height": 540,
            "slides": [
                {
                    **_slide(0, authored_shape),
                    "_yrkie_slide_width": 960,
                    "_yrkie_slide_height": 540,
                },
                {
                    **copy.deepcopy(original["slides"][0]),
                    "index": 1,
                    "_yrkie_slide_width": 720,
                    "_yrkie_slide_height": 405,
                    "_yrkie_source_slide_index": 0,
                },
            ],
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "source.pptx"
            write_pptx_from_ir(original, source_path)
            source = source_path.read_bytes()
            imported = import_pptx_to_canonical(source)["presentation"]
            desired["slides"][1] = {
                **copy.deepcopy(imported["slides"][0]),
                "index": 1,
                "_yrkie_slide_width": 720,
                "_yrkie_slide_height": 405,
                "_yrkie_source_slide_index": 0,
            }
            with patch.object(
                pptx_pipeline,
                "_normalize_editor_text_layout",
                side_effect=lambda presentation: (presentation, 1, 0),
            ):
                exported = export_ir_to_pptx(
                    desired,
                    source_pptx=source,
                    original_ir=imported,
                )

            presentation = PptxPresentation(io.BytesIO(exported))
            self.assertEqual(len(presentation.slides), 2)
            authored = presentation.slides[0].shapes[0]
            self.assertAlmostEqual(authored.width / 12700, 720, places=2)
            self.assertAlmostEqual(authored.height / 12700, 405, places=2)
            with zipfile.ZipFile(io.BytesIO(source)) as before, zipfile.ZipFile(
                io.BytesIO(exported)
            ) as after:
                self.assertEqual(
                    after.read("ppt/slides/slide1.xml"),
                    before.read("ppt/slides/slide1.xml"),
                )

    def test_source_free_scale_animation_writes_string_xml_attributes(self) -> None:
        presentation = {
            "width": 960,
            "height": 540,
            "slides": [
                {
                    **_slide(0, _shape(2, "Animated", 40)),
                    "animations": [
                        {
                            "kind": "animScale",
                            "target_shape_id": 2,
                            "scale_x": 1.25,
                            "scale_y": 0.75,
                            "duration_ms": 500,
                            "click_index": 0,
                        }
                    ],
                }
            ],
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "animated.pptx"
            write_pptx_from_ir(presentation, output_path)
            with zipfile.ZipFile(output_path) as package:
                xml = package.read("ppt/slides/slide1.xml").decode("utf-8")

        self.assertIn('<p:by x="125000" y="75000"', xml)


if __name__ == "__main__":
    unittest.main()

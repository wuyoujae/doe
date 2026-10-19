from __future__ import annotations

from test_paths import DOE_PUBLIC_ROOT

import json
import os
from pathlib import Path
import subprocess
import sys
import unittest
from modular_patch import patch
from collections import Counter
from io import BytesIO
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation

RUNTIME_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RUNTIME_ROOT))
from schema_v4_export import _render_schema_deck, _export_editable_pptx
from pptx_pipeline import parse_html_to_ir


def component_types(value):
    if isinstance(value, dict):
        if isinstance(value.get("type"), str):
            yield value["type"]
        for item in value.values():
            yield from component_types(item)
    elif isinstance(value, list):
        for item in value:
            yield from component_types(item)


class DoeComponentExportCoverageTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        result = subprocess.run(
            ["node", str(Path(__file__).with_name("doe_component_fixtures.mjs"))],
            capture_output=True, check=True, encoding="utf-8", timeout=30,
        )
        cls.fixtures = json.loads(result.stdout)

    def test_every_main_application_component_has_a_real_fixture(self):
        types = set(self.fixtures["types"])
        covered = set(component_types(self.fixtures["scenes"]))
        self.assertEqual(types - covered, set())
        self.assertNotIn("triangle-metrics", covered)
        self.assertNotIn("icon-text-card", covered)

    @unittest.skipUnless(os.environ.get("YRKIE_DOE_FULL_EXPORT_TEST") == "1", "Explicit full browser/PPTX regression")
    def test_all_components_export_as_editable_pptx(self):
        for scene in self.fixtures["scenes"]:
            with self.subTest(scene=scene["id"]):
                deck = {"version": 4, "deck": {"title": scene["id"], "theme": "yrkie-theme",
                                                "slides": [scene["schema"]["slide"]]}}
                rendered = _render_schema_deck(deck, {}, timeout_seconds=180,
                                               include_screenshots=False, include_stylesheet=True)
                snapshots = []
                def capture_layout(*args, **kwargs):
                    kwargs["include_ast"] = True
                    result = parse_html_to_ir(*args, **kwargs)
                    snapshots.append(result["ast"])
                    return result
                with patch("schema_v4_export.parse_html_to_ir", side_effect=capture_layout), \
                        patch("pptx_pipeline._LOGGER.warning") as layout_warnings:
                    payload = _export_editable_pptx(rendered["slides"], stylesheet=rendered["stylesheet"],
                                                    title=scene["id"], export_watermark=False)
                for call in layout_warnings.call_args_list:
                    self.assertNotIn("normalization failed closed", str(call.args[0]))
                    self.assertNotIn("could not be normalized", str(call.args[0]))
                pptx = Presentation(BytesIO(payload))
                self.assertEqual(len(pptx.slides), 1)
                self.assertGreater(len(pptx.slides[0].shapes), 0)
                with ZipFile(BytesIO(payload)) as package:
                    self.assertFalse(any(name.startswith("ppt/charts/") for name in package.namelist()))
                    slide = etree.fromstring(package.read("ppt/slides/slide1.xml"))
                    text = slide.xpath('//*[local-name()="t"]/text()')
                    self.assertTrue(text)
                    self.assertFalse(slide.xpath('//*[local-name()="normAutofit"]'))
                    if "table" in set(component_types(scene)):
                        self.assertTrue(slide.xpath('//*[local-name()="tbl"]'))
                    if "math-formula" in set(component_types(scene)):
                        self.assertTrue(slide.xpath('//*[local-name()="oMath"]'))
                    else:
                        visible = []
                        for node in snapshots[0]["nodes"]:
                            if node.get("nodeType") != "element" or not node.get("visible"):
                                continue
                            value = node.get("directText") or ""
                            transform = node.get("computed", {}).get("textTransform")
                            visible.append(value.upper() if transform == "uppercase" else value)
                        expected = Counter(character for character in "".join(visible) if not character.isspace())
                        actual = Counter(character for character in "".join(text) if not character.isspace())
                        self.assertFalse(expected - actual, f"Missing visible characters: {expected - actual}")
                        self.assertFalse(actual - expected, f"Duplicated or unexpected characters: {actual - expected}")


if __name__ == "__main__":
    unittest.main()

from __future__ import annotations

import base64
import copy
import dataclasses
import io
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
import modular_patch as mock

from PIL import Image
from pptx import Presentation as PptxPresentation


RUNTIME_ROOT = Path(__file__).resolve().parents[1]
if str(RUNTIME_ROOT) not in sys.path:
    sys.path.insert(0, str(RUNTIME_ROOT))

from pptx2html.opc.archive_guard import validate_pptx_archive  # noqa: E402
from pptx2html import Presentation as PomPresentation  # noqa: E402
from pptx2html.ir.builder import PresentationBuilder  # noqa: E402
from pptx2html.ir.model import (  # noqa: E402
    Border,
    Bounds,
    Fill,
    Font,
    ImageInfo,
    Paragraph,
    Presentation as IRPresentation,
    Run,
    Shape,
    Slide,
    TableCell,
    TableInfo,
    TextFrame,
)
from pptx2html.server.ooxml_pptx_writer import (  # noqa: E402
    RoundtripFallbackRequired,
    map_original_shapes_to_elements,
    write_roundtrip_pptx_from_ir,
)
from pptx2html.server.pptx_writer import write_pptx_from_ir  # noqa: E402


def _png(color: tuple[int, int, int]) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (16, 12), color).save(output, format="PNG")
    return output.getvalue()


def _picture(shape_id: int, name: str, payload: bytes) -> dict:
    return {
        "type": "picture",
        "shape_id": shape_id,
        "name": name,
        "bounds": {"left": 40, "top": 30, "width": 320, "height": 180},
        "image": {
            "data": base64.b64encode(payload).decode("ascii"),
            "content_type": "image/png",
            "crop_left": 0.0,
            "crop_right": 0.0,
            "crop_top": 0.0,
            "crop_bottom": 0.0,
        },
    }


def _presentation(*slides: dict) -> dict:
    return {"width": 960, "height": 540, "slides": list(slides)}


def _slide(index: int, picture: dict | None = None) -> dict:
    return {
        "index": index,
        "name": f"Slide {index + 1}",
        "shapes": [picture] if picture is not None else [],
    }


def _assert_valid_package(test: unittest.TestCase, path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        validate_pptx_archive(archive)
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                ET.fromstring(archive.read(name))
    PptxPresentation(path)
    test.assertGreater(path.stat().st_size, 0)


def _first_picture_color(path: Path, slide_index: int = 0) -> tuple[int, int, int]:
    presentation = PptxPresentation(path)
    picture = presentation.slides[slide_index].shapes[0]
    with Image.open(io.BytesIO(picture.image.blob)) as image:
        return image.convert("RGB").getpixel((0, 0))


class RoundtripFallbackTests(unittest.TestCase):
    def test_direct_shape_mapping_skips_inherited_id_collisions(self) -> None:
        shape_element = ET.fromstring(
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            '<p:nvSpPr><p:cNvPr id="7" name="direct"/></p:nvSpPr>'
            "</p:sp>"
        )
        mapped = map_original_shapes_to_elements(
            [
                {
                    "type": "auto_shape",
                    "shape_id": 7,
                    "name": "master",
                    "source_scope": "master",
                },
                {
                    "type": "auto_shape",
                    "shape_id": 7,
                    "name": "direct",
                    "source_scope": "slide",
                },
            ],
            [shape_element],
        )

        self.assertIsNone(mapped[0])
        self.assertIs(mapped[1], shape_element)

    def test_source_backed_edit_preserves_unedited_slide_and_roundtrips_mixed_styles(self) -> None:
        source_model = IRPresentation(
            width=960,
            height=540,
            slides=[
                Slide(
                    index=0,
                    name="Editable",
                    shapes=[
                        Shape(
                            type="auto_shape",
                            shape_id=2,
                            name="Styled text",
                            bounds=Bounds(40, 30, 380, 90),
                            geometry="rect",
                            fill=Fill(type="solid", color="F3F4F6"),
                            border=Border(color="1D4ED8", width=2),
                            text=TextFrame(
                                paragraphs=[
                                    Paragraph(
                                        runs=[
                                            Run(
                                                text="Original",
                                                font=Font(
                                                    family="Microsoft YaHei",
                                                    size_pt=24,
                                                    bold=True,
                                                    color="111827",
                                                ),
                                            )
                                        ]
                                    )
                                ]
                            ),
                        ),
                        Shape(
                            type="picture",
                            shape_id=3,
                            name="Cropped picture",
                            bounds=Bounds(460, 30, 300, 180),
                            image=ImageInfo(
                                data=base64.b64encode(_png((20, 150, 80))).decode("ascii"),
                                content_type="image/png",
                                crop_left=0.05,
                                crop_right=0.1,
                            ),
                        ),
                        Shape(
                            type="table",
                            shape_id=4,
                            name="Native table",
                            bounds=Bounds(40, 240, 600, 160),
                            table=TableInfo(
                                rows=[
                                    [
                                        TableCell(
                                            text_frame=TextFrame(
                                                paragraphs=[
                                                    Paragraph(
                                                        runs=[
                                                            Run(
                                                                text="A1",
                                                                font=Font(size_pt=16, color="FFFFFF"),
                                                            )
                                                        ]
                                                    )
                                                ]
                                            ),
                                            fill=Fill(type="solid", color="1D4ED8"),
                                        ),
                                        TableCell(text="B1"),
                                    ]
                                ],
                                col_widths=[300, 300],
                                row_heights=[160],
                            ),
                        ),
                    ],
                ),
                Slide(
                    index=1,
                    name="Untouched",
                    shapes=[
                        Shape(
                            type="auto_shape",
                            shape_id=2,
                            name="Untouched rectangle",
                            bounds=Bounds(80, 80, 240, 120),
                            geometry="roundRect",
                            fill=Fill(type="solid", color="F59E0B"),
                        )
                    ],
                ),
            ],
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_path = root / "source.pptx"
            output_path = root / "edited.pptx"
            write_pptx_from_ir(dataclasses.asdict(source_model), source_path)
            source = source_path.read_bytes()
            with PomPresentation.open(io.BytesIO(source)) as pom:
                original = dataclasses.asdict(PresentationBuilder().build(pom))
            current = copy.deepcopy(original)
            for index, slide in enumerate(current["slides"]):
                slide["_yrkie_source_slide_index"] = index

            text_shape = next(
                shape
                for shape in current["slides"][0]["shapes"]
                if shape["type"] == "auto_shape" and shape.get("text")
            )
            text_shape["bounds"]["left"] = 64
            text_shape["fill"]["color"] = "DBEAFE"
            text_shape["text"]["paragraphs"][0]["runs"][0]["text"] = "Edited"
            text_shape["text"]["paragraphs"][0]["runs"][0]["font"]["color"] = "7C3AED"

            picture = next(
                shape for shape in current["slides"][0]["shapes"] if shape["type"] == "picture"
            )
            picture["image"]["crop_left"] = 0.2
            picture["image"]["crop_bottom"] = 0.15

            table = next(
                shape for shape in current["slides"][0]["shapes"] if shape.get("table")
            )
            table["table"]["rows"][0][0]["text_frame"]["paragraphs"][0]["runs"][0]["text"] = "Edited A1"
            table["table"]["rows"][0][0]["fill"]["color"] = "0F766E"

            write_roundtrip_pptx_from_ir(
                current,
                source,
                output_path,
                original_ir=original,
            )

            _assert_valid_package(self, output_path)
            with zipfile.ZipFile(io.BytesIO(source)) as before, zipfile.ZipFile(
                output_path
            ) as after:
                self.assertEqual(
                    after.read("ppt/slides/slide2.xml"),
                    before.read("ppt/slides/slide2.xml"),
                )
                self.assertNotEqual(
                    after.read("ppt/slides/slide1.xml"),
                    before.read("ppt/slides/slide1.xml"),
                )

            with PomPresentation.open(str(output_path)) as pom:
                reparsed = dataclasses.asdict(PresentationBuilder().build(pom))
            reparsed_shapes = reparsed["slides"][0]["shapes"]
            reparsed_text = next(
                shape
                for shape in reparsed_shapes
                if shape["type"] == "auto_shape" and shape.get("text")
            )
            self.assertAlmostEqual(reparsed_text["bounds"]["left"], 64, places=2)
            self.assertEqual(reparsed_text["fill"]["color"], "DBEAFE")
            self.assertEqual(reparsed_text["text"]["paragraphs"][0]["runs"][0]["text"], "Edited")
            self.assertEqual(
                reparsed_text["text"]["paragraphs"][0]["runs"][0]["font"]["color"],
                "7C3AED",
            )
            reparsed_picture = next(
                shape for shape in reparsed_shapes if shape["type"] == "picture"
            )
            self.assertAlmostEqual(reparsed_picture["image"]["crop_left"], 0.2, places=3)
            self.assertAlmostEqual(reparsed_picture["image"]["crop_bottom"], 0.15, places=3)
            reparsed_table = next(
                shape for shape in reparsed_shapes if shape.get("table")
            )
            cell = reparsed_table["table"]["rows"][0][0]
            self.assertEqual(cell["text_frame"]["paragraphs"][0]["runs"][0]["text"], "Edited A1")
            self.assertEqual(cell["fill"]["color"], "0F766E")

    def test_unchanged_copyist_export_is_byte_identical_to_source(self) -> None:
        original = _presentation(
            _slide(0, _picture(2, "First", _png((230, 20, 30)))),
            _slide(1, _picture(2, "Second", _png((20, 70, 230)))),
        )
        current = copy.deepcopy(original)
        for index, slide in enumerate(current["slides"]):
            slide["_yrkie_source_slide_index"] = index
        current["slides"][0]["shapes"][0]["image"]["layout_snapshot"] = {"ignored": True}

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_path = root / "source.pptx"
            output_path = root / "output.pptx"
            write_pptx_from_ir(original, source_path)
            source = source_path.read_bytes()

            write_roundtrip_pptx_from_ir(
                current,
                source,
                output_path,
                original_ir=original,
            )

            self.assertEqual(output_path.read_bytes(), source)

    def test_edit_rewrites_only_the_changed_slide_part(self) -> None:
        original = _presentation(
            _slide(0, _picture(2, "First", _png((230, 20, 30)))),
            _slide(1, _picture(2, "Second", _png((20, 70, 230)))),
        )
        current = copy.deepcopy(original)
        for index, slide in enumerate(current["slides"]):
            slide["_yrkie_source_slide_index"] = index
        current["slides"][0]["shapes"][0]["bounds"]["left"] = 120

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_path = root / "source.pptx"
            output_path = root / "output.pptx"
            write_pptx_from_ir(original, source_path)
            source = source_path.read_bytes()

            write_roundtrip_pptx_from_ir(
                current,
                source,
                output_path,
                original_ir=original,
            )

            with zipfile.ZipFile(io.BytesIO(source)) as before, zipfile.ZipFile(
                output_path
            ) as after:
                self.assertEqual(
                    after.read("ppt/presentation.xml"),
                    before.read("ppt/presentation.xml"),
                )
                self.assertEqual(
                    after.read("ppt/slides/slide2.xml"),
                    before.read("ppt/slides/slide2.xml"),
                )
                self.assertNotEqual(
                    after.read("ppt/slides/slide1.xml"),
                    before.read("ppt/slides/slide1.xml"),
                )

    def test_delete_and_reorder_reject_direct_roundtrip_and_public_writer_rebuilds(self) -> None:
        red = _png((230, 20, 30))
        blue = _png((20, 70, 230))
        original = _presentation(
            _slide(0, _picture(2, "Red", red)),
            _slide(1, _picture(2, "Blue", blue)),
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_path = root / "source.pptx"
            write_pptx_from_ir(original, source_path)
            source = source_path.read_bytes()

            deleted = copy.deepcopy(original)
            deleted["slides"] = [
                {**copy.deepcopy(original["slides"][1]), "_yrkie_source_slide_index": 1}
            ]
            direct_deleted = root / "direct-deleted.pptx"
            with self.assertRaises(RoundtripFallbackRequired):
                write_roundtrip_pptx_from_ir(
                    deleted, source, direct_deleted, original_ir=original
                )
            self.assertFalse(direct_deleted.exists())

            rebuilt_deleted = root / "rebuilt-deleted.pptx"
            write_pptx_from_ir(
                deleted,
                rebuilt_deleted,
                source_pptx=source,
                original_ir=original,
            )
            _assert_valid_package(self, rebuilt_deleted)
            self.assertEqual(len(PptxPresentation(rebuilt_deleted).slides), 1)
            self.assertEqual(_first_picture_color(rebuilt_deleted), (20, 70, 230))

            reordered = copy.deepcopy(original)
            reordered["slides"] = [
                {**copy.deepcopy(original["slides"][1]), "_yrkie_source_slide_index": 1},
                {**copy.deepcopy(original["slides"][0]), "_yrkie_source_slide_index": 0},
            ]
            direct_reordered = root / "direct-reordered.pptx"
            with self.assertRaises(RoundtripFallbackRequired):
                write_roundtrip_pptx_from_ir(
                    reordered, source, direct_reordered, original_ir=original
                )
            self.assertFalse(direct_reordered.exists())

            rebuilt_reordered = root / "rebuilt-reordered.pptx"
            write_pptx_from_ir(
                reordered,
                rebuilt_reordered,
                source_pptx=source,
                original_ir=original,
            )
            _assert_valid_package(self, rebuilt_reordered)
            self.assertEqual(_first_picture_color(rebuilt_reordered, 0), (20, 70, 230))
            self.assertEqual(_first_picture_color(rebuilt_reordered, 1), (230, 20, 30))

    def test_edited_or_deleted_inherited_shape_uses_source_free_fallback(self) -> None:
        red = _png((220, 30, 20))
        green = _png((20, 210, 80))
        source_ir = _presentation(_slide(0))
        original = _presentation(_slide(0, _picture(91, "Inherited Picture", red)))
        edited = copy.deepcopy(original)
        edited_picture = edited["slides"][0]["shapes"][0]
        edited_picture["image"] = _picture(91, "Inherited Picture", green)["image"]
        edited_picture["image"]["crop_left"] = 0.2

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_path = root / "source.pptx"
            write_pptx_from_ir(source_ir, source_path)
            source = source_path.read_bytes()

            with self.assertRaises(RoundtripFallbackRequired):
                write_roundtrip_pptx_from_ir(
                    edited, source, root / "direct-edited.pptx", original_ir=original
                )

            rebuilt = root / "rebuilt-edited.pptx"
            write_pptx_from_ir(
                edited,
                rebuilt,
                source_pptx=source,
                original_ir=original,
            )
            _assert_valid_package(self, rebuilt)
            self.assertEqual(_first_picture_color(rebuilt), (20, 210, 80))
            picture = PptxPresentation(rebuilt).slides[0].shapes[0]
            self.assertAlmostEqual(picture.crop_left, 0.2, places=4)

            deleted = copy.deepcopy(original)
            deleted["slides"][0]["shapes"] = []
            with self.assertRaises(RoundtripFallbackRequired):
                write_roundtrip_pptx_from_ir(
                    deleted, source, root / "direct-deleted-shape.pptx", original_ir=original
                )
            rebuilt_deleted = root / "rebuilt-deleted-shape.pptx"
            write_pptx_from_ir(
                deleted,
                rebuilt_deleted,
                source_pptx=source,
                original_ir=original,
            )
            _assert_valid_package(self, rebuilt_deleted)
            self.assertEqual(len(PptxPresentation(rebuilt_deleted).slides[0].shapes), 0)

    def test_public_writer_does_not_swallow_unrelated_roundtrip_errors(self) -> None:
        presentation = _presentation(_slide(0))
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "output.pptx"
            with mock.patch(
                "pptx2html.server.pptx_writer.write_roundtrip_pptx_from_ir",
                side_effect=ValueError("ordinary writer failure"),
            ):
                with self.assertRaisesRegex(ValueError, "ordinary writer failure"):
                    write_pptx_from_ir(
                        presentation,
                        output_path,
                        source_pptx=b"source",
                        original_ir=presentation,
                    )
            self.assertFalse(output_path.exists())


if __name__ == "__main__":
    unittest.main()

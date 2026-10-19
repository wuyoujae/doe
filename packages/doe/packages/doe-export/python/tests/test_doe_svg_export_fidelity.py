from __future__ import annotations

from test_paths import DOE_PUBLIC_ROOT

import sys
import base64
import unittest
from io import BytesIO
from zipfile import ZipFile
from pathlib import Path

RUNTIME_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RUNTIME_ROOT))

from pptx2html.htmlir.compiler import (
    _svg_text_shape, _custom_border_radius_paths_from_node, _svg_transformed_path,
    _svg_linear_gradient_fill, _font_from_computed,
    _svg_clipped_path, _svg_stroke_dash,
    _object_fit_contain_bounds, _text_for_node, _text_frame_from_node,
    _flex_baseline_top_offset,
    _normalized_border_radius_pairs, _rounded_border_shapes_from_node,
)
from pptx2html.ir.model import Bounds
from schema_v4_native import _normalize_chart_axis_ids
from schema_v4_native import apply_schema_v4_native_components
from schema_v4_fonts import trusted_schema_v4_fonts, _apply_tabular_digits
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.util import Pt
from pptx_pipeline import parse_html_to_ir, export_ir_to_pptx


class DoeSvgExportFidelityTests(unittest.TestCase):
    def test_svg_paint_uses_resolved_theme_colors_and_css_overrides(self):
        from doe_export.html.compiler.svg._svg_paint_color import _svg_paint_color
        from doe_export.html.compiler.svg._svg_attr_style_or_computed import _svg_attr_style_or_computed

        for paint in ("fill", "stroke"):
            for attribute in ("var(--yrkie-color-primary)", "red", "currentColor"):
                node = {"attributes": {paint: attribute}, "computed": {paint: "rgb(27, 84, 218)"}}
                self.assertEqual(_svg_paint_color(node, {}, paint), ("1B54DA", 1.0))
            for value in ("none", "transparent", "url(#gradient)"):
                node = {"attributes": {paint: "blue"}, "computed": {paint: value}}
                self.assertIsNone(_svg_paint_color(node, {}, paint))
                self.assertEqual(_svg_attr_style_or_computed(node, {}, paint, paint), value)
            self.assertEqual(_svg_paint_color({"computed": {paint: "rgba(27, 84, 218, 0.5)"}}, {}, paint), ("1B54DA", .5))
            self.assertEqual(_svg_paint_color({"attributes": {paint: "currentColor"}, "computed": {"color": "rgb(27, 84, 218)"}}, {}, paint), ("1B54DA", 1.0))
            self.assertEqual(_svg_paint_color({}, {"computed": {paint: "rgb(27, 84, 218)"}}, paint), ("1B54DA", 1.0))

    def test_wave_arrow_exports_continuous_native_faces_and_white_curls(self):
        from lxml import etree
        from schema_v4_export import export_schema_v4_deck_to_pptx

        deck = {"version": 4, "deck": {"theme": "yrkie-theme", "slides": [{
            "id": "wave", "blocks": [{"id": "wave", "type": "wave-arrow", "props": {
                "items": [{"id": str(i), "title": f"Stage {i + 1}", "text": "Editable description."}
                          for i in range(3)]}}]}]}}
        data = export_schema_v4_deck_to_pptx(deck, assets={}, embed_fonts=False, timeout_seconds=120)
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
              "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
        with ZipFile(BytesIO(data)) as package:
            slide = etree.fromstring(package.read("ppt/slides/slide1.xml"))
        paints, curls = [], []
        for shape in slide.findall(".//p:sp", ns):
            geometry = shape.find(".//a:custGeom", ns)
            if geometry is None:
                continue
            fill = shape.find("./p:spPr/a:solidFill/a:srgbClr", ns)
            line = shape.find("./p:spPr/a:ln/a:solidFill/a:srgbClr", ns)
            if fill is not None and fill.get("val") in {"1B54DA", "10317F"}:
                paints.append(fill.get("val"))
                self.assertEqual(len(geometry.findall(".//a:moveTo", ns)), 1)
                self.assertEqual(len(geometry.findall(".//a:close", ns)), 1)
                self.assertGreater(len(geometry.findall(".//a:cubicBezTo", ns)), 0)
            if line is not None and line.get("val") == "FFFFFF":
                curls.append(shape)
                self.assertEqual(len(geometry.findall(".//a:close", ns)), 0)
                self.assertGreater(len(geometry.findall(".//a:cubicBezTo", ns)), 0)
        self.assertEqual(paints, ["1B54DA", "10317F", "10317F"])
        self.assertEqual(len(curls), 2)
        texts = [t.text for t in slide.findall(".//a:t", ns)]
        for expected in ["01", "02", "03", "Stage 1", "Editable description."]:
            self.assertIn(expected, texts)

    def test_anonymous_grid_text_keeps_its_measured_breaks_after_hydration(self):
        from pptx2html.ir.serde import presentation_from_dict
        from pptx2html.ir.text_layout_snapshot import layout_snapshot_is_valid, make_text_container_metrics
        result = parse_html_to_ir('''<div style="display:grid;place-items:center;width:200px;
            font:20px Alexandria;white-space:pre-wrap">Produce practical\nScenario cards</div>''', include_ast=True)
        node = next(n for n in result["ast"]["nodes"] if n.get("tag") == "div")
        self.assertIsNotNone(node.get("anonymousTextRect"))
        presentation = presentation_from_dict(result["presentation"])
        shape = next(s for s in presentation.slides[0].shapes if s.text)
        self.assertTrue(layout_snapshot_is_valid(shape.text,
            container_metrics=make_text_container_metrics(shape.bounds.width, shape.bounds.height)))
        with ZipFile(BytesIO(export_ir_to_pptx(result["presentation"]))) as package:
            self.assertEqual(package.read("ppt/slides/slide1.xml").count(b"<a:br"), 1)

    def test_aligned_text_preserves_the_browser_anchor(self):
        result = parse_html_to_ir('''<div style="display:flex;gap:8px;width:400px">
            <span style="width:24px;flex:none;text-align:right">1.</span><span>First task</span></div>
            <p style="width:200px;text-align:center">Centered heading</p>''', include_ast=True)
        shapes = result["presentation"]["slides"][0]["shapes"]
        for value in ("1.", "Centered heading"):
            node = next(n for n in result["ast"]["nodes"] if n.get("directText") == value)
            shape = next(s for s in shapes if "".join(r.get("text", "") for p in
                (s.get("text") or {}).get("paragraphs", []) for r in p["runs"]) == value)
            self.assertAlmostEqual(shape["bounds"]["left"], node["rect"]["left"] * .75, places=2)
            self.assertAlmostEqual(shape["bounds"]["width"], node["rect"]["width"] * .75, places=2)

    def test_plaintext_newlines_keep_the_original_browser_layout_snapshot(self):
        result = parse_html_to_ir('''<p style="width:600px;white-space:pre-wrap;font:18px Alexandria">
First browser line\nSecond browser line</p>'''.replace(">\nFirst", ">First"), include_ast=True)
        node = next(n for n in result["ast"]["nodes"] if n.get("tag") == "p")
        self.assertTrue(node["visualTextLayout"]["reliable"])
        self.assertEqual(len(node["visualTextLayout"]["authoredBreakOffsets"]), 1)
        shape = next(s for s in result["presentation"]["slides"][0]["shapes"] if s.get("text"))
        self.assertIsNotNone(shape["text"]["layout_snapshot"])
        self.assertEqual(shape["text"]["layout_snapshot"]["soft_breaks"], [])
        with ZipFile(BytesIO(export_ir_to_pptx(result["presentation"]))) as package:
            self.assertEqual(package.read("ppt/slides/slide1.xml").count(b"<a:br"), 1)

    def test_image_mask_preserves_transparency_without_flattening_text(self):
        from PIL import Image
        image = BytesIO()
        Image.new("RGB", (20, 20), "red").save(image, format="PNG")
        src = "data:image/png;base64," + base64.b64encode(image.getvalue()).decode()
        result = parse_html_to_ir(f'''<div style="width:100px;height:100px;
            mask-image:linear-gradient(to right,transparent 0%,black 50%)">
            <img src="{src}" style="width:100%;height:100%"></div><p>Editable text</p>''', include_ast=True)
        node = next(n for n in result["ast"]["nodes"] if n.get("tag") == "img")
        with Image.open(BytesIO(base64.b64decode(node["media"]["paintedSrc"].split(",", 1)[1]))) as masked:
            self.assertEqual(masked.size, (300, 300))
            self.assertLess(masked.getpixel((0, 150))[3], 5)
            self.assertEqual(masked.getpixel((250, 150)), (255, 0, 0, 255))
        self.assertTrue(any(s.get("text") for s in result["presentation"]["slides"][0]["shapes"]))

    def test_suppressed_pseudo_bullet_does_not_reappear_in_export(self):
        result = parse_html_to_ir('''<style>li::before {content:none;display:block;width:8px;height:8px;
            background:#000;border-radius:50%} li {list-style:none}</style>
            <ol><li><span>1.</span><span>First task</span></li></ol>''', include_ast=True)
        self.assertTrue(result["success"])
        self.assertFalse(any(node.get("tag") == "::before" for node in result["ast"]["nodes"]))

    def test_scaled_rotated_text_keeps_content_and_semantic_size(self):
        style = {"fontFamily": "Alexandria", "resolvedFontFamily": "Alexandria",
                 "fontSize": "24px", "fontWeight": "300", "fill": "rgb(23, 83, 255)"}
        node = {"directText": "DOE 2026", "computed": style,
                "svgText": {"center": {"x": 200, "y": 100}, "width": 120,
                            "height": 30, "scaleX": .5, "scaleY": .5, "rotation": 45},
                "richTextRuns": [{"kind": "text", "text": "DOE 2026", "computed": style}]}
        shape = _svg_text_shape(node, .75, .75, 1)
        self.assertEqual(shape.bounds, Bounds(105, 63.75, 90, 22.5))
        self.assertEqual(shape.rotation, 45)
        run = shape.text.paragraphs[0].runs[0]
        self.assertEqual((run.text, run.font.family, run.font.size_pt, run.font.color),
                         ("DOE 2026", "Alexandria", 9, "1753FF"))
        self.assertFalse(shape.text.word_wrap)
        self.assertEqual(shape.text.autofit_mode, "none")

    def test_parent_text_is_not_duplicated_when_tspans_own_the_lines(self):
        self.assertIsNone(_svg_text_shape({"directText": "parent", "svgText": None}, .75, .75, 1))

    def test_uniform_css_radius_is_explicit_geometry(self):
        node = {"computed": {"borderRadius": "8px"}, "rect": {"width": 1280, "height": 720}}
        paths = _custom_border_radius_paths_from_node(node, Bounds(0, 0, 960, 540))
        self.assertEqual(len(paths), 1)
        self.assertTrue(paths[0].d.startswith("M 0.833 "))
        self.assertEqual(paths[0].d.count("C "), 4)

    def test_asymmetric_radius_can_exceed_half_the_edge(self):
        node = {"computed": {"borderTopLeftRadius": "80px", "borderBottomRightRadius": "80px"},
                "layout": {"offsetWidth": 100, "offsetHeight": 100}}
        self.assertEqual(_normalized_border_radius_pairs(node, Bounds(0, 0, 100, 100)),
                         [(80, 80), (0, 0), (80, 80), (0, 0)])
        node["computed"]["borderTopRightRadius"] = "80px"
        self.assertEqual(_normalized_border_radius_pairs(node, Bounds(0, 0, 100, 100))[0], (50, 50))

    def test_different_border_widths_keep_curved_inner_and_outer_edges(self):
        import pathops
        from fontTools.svgLib.path import parse_path
        style = {"borderRadius": "8px"}
        for side in ("Top", "Right", "Bottom", "Left"):
            style.update({f"border{side}Width": "6px" if side == "Top" else "1px",
                          f"border{side}Color": "rgb(0, 0, 255)" if side == "Top" else "rgb(0, 0, 0)",
                          f"border{side}Style": "solid"})
        node = {"computed": style, "layout": {"offsetWidth": 100, "offsetHeight": 100}}
        shapes = _rounded_border_shapes_from_node(node, Bounds(10, 20, 100, 100), 1, 1, 1)
        self.assertEqual(len(shapes), 4)
        total = 0
        for shape in shapes:
            path = pathops.Path()
            parse_path(shape.geometry_paths[0].d, path.getPen())
            total += path.area
            self.assertEqual(shape.bounds, Bounds(10, 20, 100, 100))
            self.assertTrue(shape.border.no_fill)
            self.assertIn("C ", shape.geometry_paths[0].d)
        # Curved border area: outer box minus the asymmetrically inset inner box.
        self.assertAlmostEqual(total, (10000 - 4 * 64 * (1 - 3.14159265 / 4)) -
                               (98 * 93 - 2 * (7 * 2 + 7 * 7) * (1 - 3.14159265 / 4)), delta=.2)

    def test_svg_matrix_keeps_nonzero_viewbox_and_rotation(self):
        path = _svg_transformed_path("M 0 84 L 100 84", [0, 2, -2, 0, 168, 0],
                                     Bounds(0, 0, 150, 150), .75, .75)
        self.assertEqual(path, "M 0 0 L 0 100")

    def test_svg_gradient_retains_stops_and_opacity(self):
        node = {"svgGraphic": {"opacity": .5, "gradient": {
            "start": {"x": 20, "y": 0}, "end": {"x": 1020, "y": 0},
            "stops": [{"position": 0, "color": "rgb(255, 255, 255)", "alpha": 1},
                      {"position": 1, "color": "rgb(0, 0, 255)", "alpha": 1}],
        }}}
        fill = _svg_linear_gradient_fill(node, Bounds(15, 0, 750, 300), .75, .75)
        self.assertEqual(fill.gradient_angle, 0)
        self.assertEqual([(stop.position, stop.color, stop.alpha) for stop in fill.gradient_stops],
                         [(0, "FFFFFF", .5), (1, "0000FF", .5)])

    def test_tabular_digits_have_equal_advance_in_the_embedded_face(self):
        faces = trusted_schema_v4_fonts(DOE_PUBLIC_ROOT)
        face = next(face for face in faces if face.family == "Alexandria Medium Tabular")
        with TTFont(face.path) as font:
            _apply_tabular_digits(font)
            cmap = font.getBestCmap()
            widths = {font["hmtx"][cmap[ord(digit)]][0] for digit in "0123456789"}
            self.assertEqual(len(widths), 1)
        font = _font_from_computed({"fontFamily": "Alexandria Medium", "fontVariantNumeric": "tabular-nums"}, 20, 1)
        self.assertEqual(font.family, "Alexandria Medium Tabular")

    def test_native_chart_ids_preserve_references_within_office_range(self):
        source = BytesIO()
        with ZipFile(source, "w") as package:
            package.writestr("ppt/charts/chart1.xml", '<c:axId val="-2226939960"/><c:axId val="2180972856"/><c:crossAx val="-2226939960"/>')
        with ZipFile(BytesIO(_normalize_chart_axis_ids(source.getvalue()))) as package:
            self.assertEqual(package.read("ppt/charts/chart1.xml"), b'<c:axId val="1"/><c:axId val="2"/><c:crossAx val="1"/>')

    def test_vector_chart_metadata_preserves_editable_shapes(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Pt(20), Pt(20), Pt(100), Pt(30)).text = "2026 Sales"
        source = BytesIO()
        presentation.save(source)
        component = {"id": "chart", "type": "grouped-bar-chart", "representation": "editable-vector",
                     "bounds": {"left": 0, "top": 0, "width": 500, "height": 400}}
        result = apply_schema_v4_native_components(source.getvalue(), [[component]])
        restored = Presentation(BytesIO(result))
        self.assertEqual([shape.text for shape in restored.slides[0].shapes], ["2026 Sales"])
        with ZipFile(BytesIO(result)) as package:
            self.assertFalse(any(name.startswith("ppt/charts/") for name in package.namelist()))

    def test_svg_clip_intersection_preserves_curves_inside_plot(self):
        clip = {"tag": "rect", "attributes": {"x": "0", "y": "0", "width": "100", "height": "100", "rx": "8"},
                "matrix": [1, 0, 0, 1, 0, 0]}
        result = _svg_clipped_path("M -20 -20 L 120 -20 L 120 120 L -20 120 Z", [[clip]], Bounds(0, 0, 100, 100), 1, 1)
        self.assertIn("C", result)
        from fontTools.svgLib.path import parse_path
        from fontTools.pens.boundsPen import BoundsPen
        pen = BoundsPen(None)
        parse_path(result, pen)
        self.assertEqual(pen.bounds, (0, 0, 100, 100))

    def test_theme_disables_synthetic_italics_and_keeps_long_dashes(self):
        font = _font_from_computed({"fontFamily": "Alexandria", "fontStyle": "italic", "fontSynthesis": "none"}, 14, 1)
        self.assertFalse(font.italic)
        node = {"computed": {"strokeDasharray": "12px, 9px", "strokeWidth": "3px"}}
        self.assertEqual(_svg_stroke_dash(node, {}), "dash")

    def test_contained_image_preserves_aspect_and_position(self):
        node = {"computed": {"objectFit": "contain", "objectPosition": "50% 50%"},
                "media": {"naturalWidth": 1024, "naturalHeight": 1536}}
        bounds = _object_fit_contain_bounds(node, Bounds(576, 37.875, 348, 464.25))
        self.assertAlmostEqual(bounds.width / bounds.height, 2 / 3)
        self.assertAlmostEqual(bounds.left, 595.25)
        self.assertEqual(bounds.top, 37.875)

    def test_anonymous_grid_text_has_no_padding_or_added_wrap(self):
        node = {"tag": "article", "path": "article", "directText": "Revenue", "visualTextRaw": "Revenue",
                "anonymousTextRect": {"width": 90, "lineCount": 1},
                "computed": {"display": "grid", "paddingTop": "18px", "paddingLeft": "20px", "fontSize": "20px"}}
        self.assertEqual(_text_for_node(node, {"article", "article > svg"}, {}), "Revenue")
        frame = _text_frame_from_node("Revenue", node, .75, .75)
        self.assertFalse(frame.word_wrap)
        self.assertEqual(frame.margins, {"left": 0, "right": 0, "top": 0, "bottom": 0})

    def test_absent_table_border_overrides_office_theme(self):
        from pptx2html.server.pptx_writer import _table_cell_border_xml
        border = _table_cell_border_xml("a:lnL", None)
        self.assertEqual(border[0].tag.rsplit("}", 1)[-1], "noFill")

    def test_doe_multiline_baseline_uses_measured_glyph_overflow(self):
        node = {"path": "section > h1", "computed": {"fontFamily": "MiSans"},
                "rect": {"top": 50}, "textRect": {"top": 45},
                "visualTextLayout": {"lines": [{}]}}
        parents = {"section": {"computed": {"display": "block"}}}
        self.assertEqual(_flex_baseline_top_offset(node, parents, .75), 3.75)
        node["visualTextLayout"]["lines"].append({})
        self.assertEqual(_flex_baseline_top_offset(node, parents, .75), 3.75)

    def test_browser_keeps_zero_height_gridlines_and_anonymous_legend_text(self):
        html = '''<html><body style="margin:0"><section class="slide" style="width:1280px;height:720px">
          <svg width="400" height="200" viewBox="0 0 100 50">
            <defs><clipPath id="clip"><rect width="100" height="50" rx="3"/></clipPath>
              <marker id="arrow" markerWidth="8" markerHeight="8" refX="6" refY="3" orient="auto"><path d="M0 0 L0 6 L6 3Z"/></marker></defs>
            <g clip-path="url(#clip)"><ellipse cx="100" cy="0" rx="100" ry="50" fill="blue"/></g>
            <line x1="0" y1="25" x2="100" y2="25" stroke="blue" stroke-width="1" style="vector-effect:non-scaling-stroke"/>
            <line x1="10" y1="40" x2="90" y2="40" stroke="blue" marker-end="url(#arrow)"/>
            <text x="10" y="20" fill="blue" font-size="4">2026</text>
          </svg>
          <div style="display:flex;align-items:center;gap:10px;font:14px Arial;width:150px;height:20px"><i style="display:block;width:30px;height:3px;background:blue"></i>Revenue</div>
          <h1 style="font-family:Alexandria,MiSans;font-weight:400">GIF 动图 · 原生</h1>
        </section></body></html>'''
        result = parse_html_to_ir(html, include_ast=True, root_selector=".slide", viewport=(1280, 720),
                                  timeout_ms=60000, render_mode="semantic")
        self.assertTrue(result["success"])
        nodes = result["ast"]["nodes"]
        line = next(node for node in nodes if node.get("tag") == "line")
        self.assertTrue(line["visible"])
        self.assertTrue(line["svgGraphic"]["nonScalingStroke"])
        ellipse = next(node for node in nodes if node.get("tag") == "ellipse")
        self.assertEqual(len(ellipse["svgGraphic"]["clips"]), 1)
        arrow = next(node for node in nodes if node.get("tag") == "line" and node["svgGraphic"]["markers"])
        self.assertEqual(len(arrow["svgGraphic"]["markers"]), 1)
        legend = next(node for node in nodes if node.get("directText") == "Revenue")
        self.assertAlmostEqual(legend["anonymousTextRect"]["left"], 40, places=2)
        heading = next(node for node in nodes if node.get("tag") == "h1")
        runs = heading["richTextRuns"]
        self.assertEqual([(run["text"], run["computed"]["resolvedFontFamily"]) for run in runs],
                         [("GIF ", "Alexandria Medium"), ("动图", "MiSans Medium"),
                          (" · ", "Alexandria Medium"), ("原生", "MiSans Medium")])


if __name__ == "__main__":
    unittest.main()

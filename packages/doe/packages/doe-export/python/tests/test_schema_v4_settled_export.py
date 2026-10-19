"""Image metadata must finish reflowing before PDF/PPTX snapshots are frozen."""
import base64
import copy
from io import BytesIO
import re
import shutil
import unittest

import fitz
from PIL import Image
from pptx import Presentation
from schema_v4_export import _render_schema_deck, export_schema_v4_deck_to_pptx
from test_paths import DOE_PUBLIC_ROOT


def process_with_image():
    image = BytesIO()
    Image.new('RGB', (1024, 768), '#346954').save(image, format='PNG')
    uri = 'data:image/png;base64,' + base64.b64encode(image.getvalue()).decode()
    items = [
        {'title': 'Discover needs', 'text': 'Discuss the space, light and daily routine.', 'icon': 'Search'},
        {'title': 'Prepare materials', 'text': 'Choose materials based on the discussion.', 'icon': 'Package'},
        {'title': 'Deliver a session', 'text': 'Introduce the materials and practise together.', 'icon': 'Users'},
        {'title': 'Gather feedback', 'text': 'Record what worked and what should change.', 'icon': 'ClipboardCheck'},
    ]
    deck = {'version': 4, 'deck': {'theme': 'yrkie-botanical', 'slides': [{
        'id': 'metadata-process', 'blocks': [{
            'id': 'row', 'type': 'layout-group',
            'layout': {'variant': 'composition', 'ratio': 'primary-first', 'direction': 'row'},
            'children': [
                {'id': 'copy', 'type': 'layout-group', 'layout': {'variant': 'stack'}, 'children': [
                    {'id': 'title', 'type': 'heading', 'props': {'text': 'A four-stage experience', 'level': 1, 'variant': 'plain'}},
                    {'id': 'intro', 'type': 'text', 'props': {'text': 'Four stages, in their original reading order.'}},
                    {'id': 'steps', 'type': 'contents', 'props': {'presentation': 'numbered-parallelogram',
                        'orientation': 'vertical', 'staggered': True, 'connectors': True, 'items': items}},
                ]},
                {'id': 'image', 'type': 'media', 'props': {'src': uri, 'fit': 'cover', 'aspectRatio': '4:3'}},
            ],
        }],
    }]}}
    return deck, items


@unittest.skipUnless(shutil.which('node') and DOE_PUBLIC_ROOT.is_dir(), 'built DOE runtime required')
class SettledSchemaExportTests(unittest.TestCase):
    def test_image_reflow_keeps_all_four_stages_in_pdf_and_editable_pptx(self):
        deck, items = process_with_image()
        original = copy.deepcopy(deck)
        result = _render_schema_deck(deck, {}, timeout_seconds=180,
            include_screenshots=False, include_stylesheet=True, output_format='pdf')
        with fitz.open(stream=base64.b64decode(result['pdf_base64']), filetype='pdf') as pdf:
            self.assertEqual(len(pdf), 1)
            self.assertTrue(pdf[0].get_fonts(), 'A blank snapshot cannot pass as a successful PDF')
            self.assertTrue(pdf[0].get_drawings())
            text = re.sub(r'\s+', '', pdf[0].get_text())
            for item in items:
                for field in ('title', 'text'):
                    self.assertIn(re.sub(r'\s+', '', item[field]), text)
        output = export_schema_v4_deck_to_pptx(deck, embed_fonts=False, timeout_seconds=180)
        pptx = Presentation(BytesIO(output))
        self.assertEqual(len(pptx.slides), 1)
        text_shapes = [shape for shape in pptx.slides[0].shapes if shape.has_text_frame]
        self.assertGreaterEqual(len(text_shapes), 10, 'Content must remain native editable text')
        text = re.sub(r'\s+', '', ''.join(shape.text for shape in text_shapes))
        for item in items:
            for field in ('title', 'text'):
                self.assertIn(re.sub(r'\s+', '', item[field]), text)
        self.assertEqual(deck, original, 'Export must not change author content')

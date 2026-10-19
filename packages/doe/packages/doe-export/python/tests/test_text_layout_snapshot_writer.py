from __future__ import annotations

import copy
import dataclasses
import unittest
import xml.etree.ElementTree as ET

from pptx import Presentation as PptxPresentation
from pptx.util import Pt

from pptx2html.ir.text_layout_snapshot import (
    InvalidTextLayoutSnapshot,
    NO_WRAP_LAYOUT_SNAPSHOT_PROVIDER,
    create_layout_snapshot,
    layout_snapshot_is_valid,
    make_text_container_metrics,
    materialize_layout_snapshot,
    refresh_no_wrap_layout_snapshot,
)
from pptx2html.server import ooxml_pptx_writer, pptx_writer


A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _frame(*, hard_break: bool = False) -> dict:
    runs = [
        {
            "kind": "text",
            "text": "Alpha😀Beta",
            "font": {
                "family": "Noto Sans SC",
                "size_pt": 24.0,
                "bold": True,
                "color": "112233",
            },
        }
    ]
    if hard_break:
        runs.extend(
            [
                {"kind": "break", "text": "", "font": {"size_pt": 24.0}},
                {
                    "kind": "text",
                    "text": "Gamma",
                    "font": {"family": "Noto Sans SC", "size_pt": 18.0},
                },
            ]
        )
    return {
        "vertical_align": "top",
        "word_wrap": True,
        "white_space": "normal",
        "overflow_wrap": "normal",
        "word_break": "normal",
        "margins": {"left": 0.0, "right": 0.0, "top": 0.0, "bottom": 0.0},
        "autofit_mode": "none",
        "paragraphs": [
            {
                "alignment": "l",
                "line_spacing_pt": 30.0,
                "line_spacing_type": "points",
                "runs": runs,
            }
        ],
    }


def _with_snapshot(
    frame: dict,
    *,
    offset_utf16: int = 7,
    metrics: dict | None = None,
) -> tuple[dict, dict]:
    metrics = metrics or make_text_container_metrics(240.0, 90.0)
    snapshot = create_layout_snapshot(
        frame,
        [{"paragraph_index": 0, "run_index": 0, "offset_utf16": offset_utf16}],
        container_metrics=metrics,
    )
    result = copy.deepcopy(frame)
    result["layout_snapshot"] = dataclasses.asdict(snapshot)
    result["layout_snapshot_dirty"] = False
    return result, metrics


def _python_writer_xml(frame: dict, metrics: dict) -> ET.Element:
    presentation = PptxPresentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Pt(0), Pt(0), Pt(240), Pt(90))
    pptx_writer._apply_text_frame(
        shape.text_frame,
        frame,
        container_metrics=metrics,
    )
    return ET.fromstring(shape.text_frame._txBody.xml.encode("utf-8"))


def _native_writer_xml(frame: dict, metrics: dict) -> ET.Element:
    return ooxml_pptx_writer.build_text_body(frame, container_metrics=metrics)


def _assert_writer_layout(
    testcase: unittest.TestCase,
    root: ET.Element,
    *,
    wrap: str,
    breaks: int,
) -> None:
    body_pr = root.find(f"{{{A}}}bodyPr")
    testcase.assertIsNotNone(body_pr)
    testcase.assertEqual(body_pr.get("wrap"), wrap)
    testcase.assertEqual(len(root.findall(f".//{{{A}}}br")), breaks)


class TextLayoutSnapshotWriterTests(unittest.TestCase):
    def test_baseline_top_offset_is_added_to_the_top_inset(self) -> None:
        frame = _frame()
        frame["margins"]["top"] = 1.5
        frame["baseline_top_offset"] = 4.5
        metrics = make_text_container_metrics(240.0, 90.0)

        for writer in (_python_writer_xml, _native_writer_xml):
            with self.subTest(writer=writer.__name__):
                root = writer(frame, metrics)
                body_pr = root.find(f"{{{A}}}bodyPr")
                self.assertIsNotNone(body_pr)
                self.assertEqual(body_pr.get("tIns"), str(round(6.0 * 12_700)))

    def test_explicit_no_wrap_frame_uses_deterministic_empty_snapshot(self) -> None:
        frame = _frame()
        frame["word_wrap"] = False
        metrics = make_text_container_metrics(240.0, 90.0)

        self.assertTrue(
            refresh_no_wrap_layout_snapshot(
                frame,
                container_metrics=metrics,
            )
        )
        self.assertEqual(
            frame["layout_snapshot"]["provider"],
            NO_WRAP_LAYOUT_SNAPSHOT_PROVIDER,
        )
        self.assertEqual(frame["layout_snapshot"]["soft_breaks"], [])
        self.assertFalse(frame["layout_snapshot_dirty"])
        self.assertTrue(layout_snapshot_is_valid(frame, container_metrics=metrics))
        for writer in (_python_writer_xml, _native_writer_xml):
            with self.subTest(writer=writer.__name__):
                _assert_writer_layout(self, writer(frame, metrics), wrap="none", breaks=0)

    def test_deterministic_no_wrap_provider_rejects_wrappable_text(self) -> None:
        with self.assertRaises(InvalidTextLayoutSnapshot):
            create_layout_snapshot(
                _frame(),
                [],
                container_metrics=make_text_container_metrics(240.0, 90.0),
                provider=NO_WRAP_LAYOUT_SNAPSHOT_PROVIDER,
            )

    def test_valid_snapshot_is_consumed_by_both_writers_without_mutating_ir(self) -> None:
        frame, metrics = _with_snapshot(_frame())
        original = copy.deepcopy(frame)

        for writer in (_python_writer_xml, _native_writer_xml):
            with self.subTest(writer=writer.__name__):
                root = writer(frame, metrics)
                _assert_writer_layout(self, root, wrap="none", breaks=1)

        self.assertEqual(frame, original)
        materialized = materialize_layout_snapshot(frame, container_metrics=metrics)
        self.assertIsNotNone(materialized)
        self.assertFalse(materialized["word_wrap"])

    def test_stale_source_is_ignored_by_both_writers(self) -> None:
        frame, metrics = _with_snapshot(_frame())
        frame["paragraphs"][0]["runs"][0]["text"] = "Edited after capture"

        self.assertFalse(layout_snapshot_is_valid(frame, container_metrics=metrics))
        for writer in (_python_writer_xml, _native_writer_xml):
            with self.subTest(writer=writer.__name__):
                _assert_writer_layout(self, writer(frame, metrics), wrap="square", breaks=0)

    def test_stale_metrics_and_dirty_flag_are_ignored(self) -> None:
        frame, metrics = _with_snapshot(_frame())
        changed_metrics = make_text_container_metrics(241.0, 90.0)
        self.assertFalse(layout_snapshot_is_valid(frame, container_metrics=changed_metrics))
        _assert_writer_layout(
            self,
            _native_writer_xml(frame, changed_metrics),
            wrap="square",
            breaks=0,
        )

        frame["layout_snapshot_dirty"] = True
        self.assertFalse(layout_snapshot_is_valid(frame, container_metrics=metrics))
        _assert_writer_layout(
            self,
            _python_writer_xml(frame, metrics),
            wrap="square",
            breaks=0,
        )

    def test_non_metric_alignment_and_color_edits_preserve_snapshot(self) -> None:
        frame, metrics = _with_snapshot(_frame())
        frame["paragraphs"][0]["alignment"] = "r"
        frame["paragraphs"][0]["runs"][0]["font"]["color"] = "FF0000"

        self.assertTrue(layout_snapshot_is_valid(frame, container_metrics=metrics))
        _assert_writer_layout(
            self,
            _native_writer_xml(frame, metrics),
            wrap="none",
            breaks=1,
        )

    def test_utf16_surrogate_offsets_split_only_at_codepoint_boundaries(self) -> None:
        frame = _frame()
        frame["paragraphs"][0]["runs"][0]["text"] = "A😀BC"
        snapshotted, metrics = _with_snapshot(frame, offset_utf16=3)
        materialized = materialize_layout_snapshot(snapshotted, container_metrics=metrics)

        self.assertEqual(
            [
                run.get("text")
                for run in materialized["paragraphs"][0]["runs"]
                if run.get("kind") != "break"
            ],
            ["A😀", "BC"],
        )
        with self.assertRaises(InvalidTextLayoutSnapshot):
            _with_snapshot(frame, offset_utf16=2)

    def test_soft_break_coordinates_may_not_target_paragraph_endpoints(self) -> None:
        frame = _frame()
        metrics = make_text_container_metrics(240.0, 90.0)
        utf16_length = len("Alpha😀Beta".encode("utf-16-le")) // 2
        for offset in (0, utf16_length):
            with self.subTest(offset=offset), self.assertRaises(InvalidTextLayoutSnapshot):
                create_layout_snapshot(
                    frame,
                    [{"paragraph_index": 0, "run_index": 0, "offset_utf16": offset}],
                    container_metrics=metrics,
                )

    def test_rich_run_boundary_is_canonicalized_to_next_run_and_deduplicated(self) -> None:
        frame = _frame()
        frame["paragraphs"][0]["runs"] = [
            {"kind": "text", "text": "Alpha", "font": {"size_pt": 20.0}},
            {"kind": "text", "text": "Beta", "font": {"size_pt": 20.0, "bold": True}},
        ]
        metrics = make_text_container_metrics(240.0, 90.0)
        snapshot = create_layout_snapshot(
            frame,
            [
                {"paragraph_index": 0, "run_index": 0, "offset_utf16": 5},
                {"paragraph_index": 0, "run_index": 1, "offset_utf16": 0},
            ],
            container_metrics=metrics,
        )
        frame["layout_snapshot"] = dataclasses.asdict(snapshot)
        materialized = materialize_layout_snapshot(frame, container_metrics=metrics)

        self.assertEqual(len(snapshot.soft_breaks), 1)
        self.assertEqual(dataclasses.asdict(snapshot.soft_breaks[0]), {
            "paragraph_index": 0,
            "run_index": 1,
            "offset_utf16": 0,
        })
        self.assertEqual(
            [run.get("kind", "text") for run in materialized["paragraphs"][0]["runs"]],
            ["text", "break", "text"],
        )

    def test_field_soft_break_preserves_field_semantics_and_style(self) -> None:
        frame = _frame()
        frame["paragraphs"][0]["runs"] = [
            {
                "kind": "field",
                "text": "Displayed page number 17",
                "field_id": "field-17",
                "field_type": "slidenum",
                "font": {"family": "Arial", "size_pt": 20.0, "bold": True},
            }
        ]
        metrics = make_text_container_metrics(240.0, 90.0)
        snapshot = create_layout_snapshot(
            frame,
            [{"paragraph_index": 0, "run_index": 0, "offset_utf16": 10}],
            container_metrics=metrics,
        )
        frame["layout_snapshot"] = dataclasses.asdict(snapshot)
        materialized = materialize_layout_snapshot(frame, container_metrics=metrics)

        runs = materialized["paragraphs"][0]["runs"]
        self.assertEqual([run["kind"] for run in runs], ["field", "break", "field"])
        self.assertEqual(
            [run.get("field_id") for run in runs if run["kind"] == "field"],
            ["field-17", "field-17"],
        )
        self.assertEqual(
            [run.get("field_type") for run in runs if run["kind"] == "field"],
            ["slidenum", "slidenum"],
        )
        self.assertTrue(all(run["font"]["bold"] for run in runs if run["kind"] == "field"))

    def test_soft_break_cannot_duplicate_an_authored_break_boundary(self) -> None:
        frame = _frame(hard_break=True)
        metrics = make_text_container_metrics(240.0, 90.0)
        utf16_length = len("Alpha😀Beta".encode("utf-16-le")) // 2

        with self.assertRaises(InvalidTextLayoutSnapshot):
            create_layout_snapshot(
                frame,
                [
                    {
                        "paragraph_index": 0,
                        "run_index": 0,
                        "offset_utf16": utf16_length,
                    }
                ],
                container_metrics=metrics,
            )

    def test_authored_and_soft_breaks_remain_distinct_and_ordered(self) -> None:
        frame, metrics = _with_snapshot(_frame(hard_break=True))
        materialized = materialize_layout_snapshot(frame, container_metrics=metrics)
        kinds = [run.get("kind", "text") for run in materialized["paragraphs"][0]["runs"]]

        self.assertEqual(kinds, ["text", "break", "text", "break", "text"])
        for writer in (_python_writer_xml, _native_writer_xml):
            with self.subTest(writer=writer.__name__):
                _assert_writer_layout(self, writer(frame, metrics), wrap="none", breaks=2)

    def test_htmlir_collapsible_spaces_are_trimmed_at_export_line_edges(self) -> None:
        frame = _frame(hard_break=True)
        frame["style_provenance"] = ["htmlir:computed-style"]
        frame["paragraphs"][0]["runs"][2]["text"] = " Our 2026 Journey "
        frame, metrics = _with_snapshot(frame)
        original = copy.deepcopy(frame)

        for writer in (_python_writer_xml, _native_writer_xml):
            with self.subTest(writer=writer.__name__):
                root = writer(frame, metrics)
                texts = [node.text or "" for node in root.findall(f".//{{{A}}}t")]
                self.assertEqual(texts[-1], "Our 2026 Journey")
                _assert_writer_layout(self, root, wrap="none", breaks=2)

        self.assertEqual(frame, original)

    def test_htmlir_soft_break_keeps_its_semantic_word_separator(self) -> None:
        frame = _frame()
        frame["style_provenance"] = ["htmlir:computed-style"]
        frame["paragraphs"][0]["runs"][0]["text"] = "Alpha Beta"
        frame, metrics = _with_snapshot(frame, offset_utf16=6)

        for writer in (_python_writer_xml, _native_writer_xml):
            with self.subTest(writer=writer.__name__):
                root = writer(frame, metrics)
                texts = [node.text or "" for node in root.findall(f".//{{{A}}}t")]
                self.assertEqual("".join(texts), "Alpha Beta")
                self.assertTrue(texts[0].endswith(" "))

    def test_htmlir_preserving_white_space_keeps_post_break_indent(self) -> None:
        frame = _frame(hard_break=True)
        frame["style_provenance"] = ["htmlir:computed-style"]
        frame["white_space"] = "pre-wrap"
        frame["paragraphs"][0]["runs"][2]["text"] = "  Second"
        metrics = make_text_container_metrics(240.0, 90.0)

        for writer in (_python_writer_xml, _native_writer_xml):
            with self.subTest(writer=writer.__name__):
                root = writer(frame, metrics)
                texts = [node.text or "" for node in root.findall(f".//{{{A}}}t")]
                self.assertEqual(texts[-1], "  Second")

    def test_imported_ir_without_snapshot_keeps_legacy_wrap_and_hard_break(self) -> None:
        frame = _frame(hard_break=True)
        metrics = make_text_container_metrics(240.0, 90.0)
        for writer in (_python_writer_xml, _native_writer_xml):
            with self.subTest(writer=writer.__name__):
                _assert_writer_layout(self, writer(frame, metrics), wrap="square", breaks=1)

    def test_runtime_snapshot_metadata_does_not_mark_source_shape_edited(self) -> None:
        original = _frame()
        original["diagnostics"] = []
        current, _metrics = _with_snapshot(original)
        current["layout_snapshot_dirty"] = True
        current["diagnostics"].append("htmlir:visual-layout-snapshot")

        self.assertFalse(ooxml_pptx_writer.shape_changed(current, original))
        current["paragraphs"][0]["runs"][0]["text"] = "Real edit"
        self.assertTrue(ooxml_pptx_writer.shape_changed(current, original))

    def test_snapshot_diff_filter_retains_real_diagnostics_and_style_changes(self) -> None:
        original = _frame()
        original["diagnostics"] = []
        original["style_provenance"] = ["pptx:body"]
        current, _metrics = _with_snapshot(original)
        current["diagnostics"] = ["htmlir:visual-layout-snapshot", "htmlir:real-warning"]

        self.assertTrue(ooxml_pptx_writer.shape_changed(current, original))
        current["diagnostics"] = ["htmlir:visual-layout-snapshot"]
        current["style_provenance"] = ["pptx:body", "htmlir:computed-style"]
        self.assertTrue(ooxml_pptx_writer.shape_changed(current, original))

    def test_snapshot_diagnostic_inside_table_cell_does_not_mark_table_edited(self) -> None:
        original_frame = _frame()
        original_frame["diagnostics"] = []
        current_frame, _metrics = _with_snapshot(original_frame)
        current_frame["diagnostics"] = ["htmlir:visual-layout-snapshot"]
        original_table = {
            "col_widths": [240.0],
            "row_heights": [90.0],
            "rows": [[{"text_frame": original_frame, "margins": {"left": 0.0, "right": 0.0}}]],
        }
        current_table = copy.deepcopy(original_table)
        current_table["rows"][0][0]["text_frame"] = current_frame

        self.assertFalse(ooxml_pptx_writer.shape_changed(current_table, original_table))
        current_table["rows"][0][0]["text_frame"]["diagnostics"].append(
            "htmlir:table-real-warning"
        )
        self.assertTrue(ooxml_pptx_writer.shape_changed(current_table, original_table))

    def test_source_backed_patch_preserves_txbody_for_snapshot_only_enrichment(self) -> None:
        original_text = _frame()
        original_text["diagnostics"] = []
        current_text, _metrics = _with_snapshot(original_text)
        current_text["diagnostics"] = ["htmlir:visual-layout-snapshot"]
        original_shape = {
            "type": "auto_shape",
            "shape_id": 2,
            "name": "TextBox 1",
            "text": original_text,
        }
        current_shape = {**original_shape, "text": current_text}
        source = ET.fromstring(
            f"""<p:sp xmlns:p="{ooxml_pptx_writer.P}" xmlns:a="{A}">
              <p:spPr/>
              <p:txBody><a:bodyPr upright="1"/><a:lstStyle/><a:p><a:r><a:t>Alpha😀Beta</a:t></a:r></a:p></p:txBody>
            </p:sp>"""
        )

        patched = ooxml_pptx_writer.patch_shape_element(
            None,
            source,
            current_shape,
            original_shape,
        )

        body_pr = patched.find(f"{{{ooxml_pptx_writer.P}}}txBody/{{{A}}}bodyPr")
        self.assertEqual(body_pr.get("upright"), "1")

    def test_table_cell_metrics_include_span_and_cell_margins(self) -> None:
        margins = {"left": 5.0, "right": 7.0, "top": 3.0, "bottom": 4.0}
        cell = {"colspan": 2, "rowspan": 2, "margins": margins}
        actual = ooxml_pptx_writer.table_cell_text_container_metrics(
            cell,
            col_index=1,
            row_index=0,
            col_widths=[40.0, 60.0, 80.0],
            row_heights=[30.0, 45.0],
        )

        self.assertEqual(
            actual,
            make_text_container_metrics(
                140.0,
                75.0,
                kind="table-cell",
                cell_margins=margins,
            ),
        )

    def test_custom_text_rectangle_drives_snapshot_metrics(self) -> None:
        shape = {
            "bounds": {"left": 0.0, "top": 0.0, "width": 200.0, "height": 100.0},
            "text_rect": {"left": 25.0, "top": 10.0, "right": 75.0, "bottom": 90.0},
        }
        self.assertEqual(
            ooxml_pptx_writer.shape_text_container_metrics(shape),
            make_text_container_metrics(100.0, 80.0),
        )
        self.assertEqual(
            pptx_writer._shape_text_container_metrics(shape, 200.0, 100.0),
            make_text_container_metrics(100.0, 80.0),
        )


if __name__ == "__main__":
    unittest.main()

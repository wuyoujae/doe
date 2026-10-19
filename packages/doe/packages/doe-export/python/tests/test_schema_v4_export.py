from __future__ import annotations

from test_paths import DOE_PUBLIC_ROOT

import os
import shutil
import subprocess
import sys
import unittest
from io import BytesIO
from pathlib import Path
from modular_patch import patch
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.enum.chart import XL_LEGEND_POSITION
from pptx.util import Pt


RUNTIME_ROOT = Path(__file__).resolve().parents[1]
if str(RUNTIME_ROOT) not in sys.path:
    sys.path.insert(0, str(RUNTIME_ROOT))

from schema_v4_export import (  # noqa: E402
    COMPILER_CONTRACT,
    RENDERER_CONTRACT,
    WRITER_CONTRACT,
    SchemaV4ExportError,
    _merge_presentations,
    _render_schema_deck,
    _static_slide_html,
    _validate_deck,
    export_schema_v4_deck_to_pptx,
    render_schema_v4_thumbnail,
)
from schema_v4_math import MATH_NS, _mathml_to_omml, formula_spec  # noqa: E402
from schema_v4_native import (  # noqa: E402
    _component_font_name,
    apply_schema_v4_native_components,
)
from schema_v4_fonts import (  # noqa: E402
    PreparedFontFace,
    _collect_visible_characters,
    _embed_font_parts,
    trusted_schema_v4_fonts,
)
from pptx2html.htmlir import compiler as compiler_module  # noqa: E402
from pptx_pipeline import parse_html_to_ir as pipeline_parse_html_to_ir  # noqa: E402


class SchemaV4ExportTests(unittest.TestCase):
    @staticmethod
    def _screenshot_base64(color: tuple[int, int, int] = (24, 67, 109)) -> str:
        output = BytesIO()
        Image.new("RGB", (1280, 720), color).save(output, format="PNG")
        import base64

        return base64.b64encode(output.getvalue()).decode("ascii")

    def test_requires_version_four_and_non_empty_slides(self) -> None:
        with self.assertRaisesRegex(SchemaV4ExportError, "version 4"):
            _validate_deck({"version": 3, "deck": {"slides": [{}]}})
        with self.assertRaisesRegex(SchemaV4ExportError, "between 1 and"):
            _validate_deck({"version": 4, "deck": {"slides": []}})
        self.assertEqual(
            _validate_deck({"version": 4, "deck": {"slides": [{"id": "s1"}]}}),
            [{"id": "s1"}],
        )

    def test_static_html_disables_motion_and_preserves_root(self) -> None:
        output = _static_slide_html(
            ".schema-slide{color:#123456}",
            '<section class="schema-slide"><p>hello</p></section>',
        )
        self.assertIn("animation:none!important", output)
        self.assertIn("transition:none!important", output)
        self.assertIn('<section class="schema-slide">', output)
        self.assertIn(".schema-slide{color:#123456}", output)

    def test_schema_renderer_removes_font_face_rules_before_resource_inlining(self) -> None:
        node = os.environ.get("YRKIE_HTMLIR_NODE_EXECUTABLE") or shutil.which("node")
        if not node:
            self.skipTest("Node.js is unavailable")
        renderer = RUNTIME_ROOT / "schema_v4_renderer.js"
        script = (
            "const fs=require('fs');"
            "const {stripFontFaceRules}=require(process.argv[1]);"
            "process.stdout.write(stripFontFaceRules(fs.readFileSync(0,'utf8')));"
        )
        source = (
            '@font-face{font-family:"Deck";src:url("missing.woff2")}\n'
            ".schema-slide{font-family:\"Deck\";color:#123456}\n"
            "@media print{@font-face{font-family:\"Nested\";src:url(\"nested.woff2\")}"
            ".schema-slide{color:#654321}}"
        )
        completed = subprocess.run(
            [node, "-e", script, str(renderer)],
            input=source,
            capture_output=True,
            text=True,
            timeout=30,
            check=True,
        )
        self.assertNotIn("@font-face", completed.stdout.lower())
        self.assertIn(".schema-slide", completed.stdout)
        self.assertIn("font-family:\"Deck\"", completed.stdout)

    def test_pipeline_forwards_the_schema_snapshot_root_selector(self) -> None:
        with patch(
            "pptx_pipeline.compile_html",
            return_value={"presentation": {"slides": []}},
        ) as compile_mock:
            pipeline_parse_html_to_ir(
                '<section class="schema-slide"></section>',
                root_selector=".schema-slide",
            )
        self.assertEqual(
            compile_mock.call_args.kwargs["root_selector"],
            ".schema-slide",
        )

    def test_browser_snapshot_payload_uses_the_requested_root_selector(self) -> None:
        with patch.object(
            compiler_module,
            "_invoke_browser_snapshot",
            return_value={"success": True},
        ) as invoke_mock:
            compiler_module._run_browser_snapshot(
                source='<section class="schema-slide"></section>',
                base_url=None,
                root_selector=".schema-slide",
                viewport=(1280, 720),
                timeout_ms=20_000,
                max_dom_nodes=100,
            )
        self.assertEqual(
            invoke_mock.call_args.args[0]["root_selector"],
            ".schema-slide",
        )

    def test_export_runtime_applies_the_deck_font_to_the_schema_root(self) -> None:
        runtime_html = (
            DOE_PUBLIC_ROOT
            / "export-runtime"
            / "index.html"
        ).read_text(encoding="utf-8")
        self.assertIn("#export-slide-root > .schema-slide", runtime_html)
        self.assertIn("font-family: var(--font-sans-cn)", runtime_html)
        self.assertIn("font-synthesis: none", runtime_html)

    def test_merge_requires_exactly_one_ir_slide_per_compiled_slide(self) -> None:
        merged = _merge_presentations(
            [
                {"name": "first", "slides": [{"name": "s1"}]},
                {"name": "second", "slides": [{"name": "s2"}]},
            ],
            title="Deck title",
        )
        self.assertEqual(merged["name"], "Deck title")
        self.assertEqual([slide["name"] for slide in merged["slides"]], ["s1", "s2"])
        with self.assertRaisesRegex(SchemaV4ExportError, "exactly one"):
            _merge_presentations([{"slides": []}], title="invalid")

    def test_visual_pipeline_places_exact_browser_screenshot_on_each_slide(self) -> None:
        deck = {
            "version": 4,
            "deck": {
                "title": "visual test",
                "slides": [{"id": "s1"}, {"id": "s2"}],
            },
        }
        rendered = {
            "success": True,
            "renderer_contract": RENDERER_CONTRACT,
            "stylesheet": "",
            "slides": [
                {"screenshot_png_base64": self._screenshot_base64((24, 67, 109))},
                {"screenshot_png_base64": self._screenshot_base64((221, 70, 62))},
            ],
        }
        with (
            patch("schema_v4_export._render_schema_deck", return_value=rendered) as render_mock,
            patch("schema_v4_export.parse_html_to_ir") as compile_mock,
            patch("schema_v4_export.export_ir_to_pptx") as writer_mock,
            patch("schema_v4_export.apply_pptx_export_watermark", side_effect=lambda value: value) as watermark_mock,
        ):
            output = export_schema_v4_deck_to_pptx(
                deck,
                export_watermark=True,
                output_mode="visual",
            )

        exported = Presentation(BytesIO(output))
        self.assertEqual(len(exported.slides), 2)
        self.assertEqual([len(slide.shapes) for slide in exported.slides], [1, 1])
        for slide in exported.slides:
            picture = slide.shapes[0]
            self.assertEqual((picture.left, picture.top), (0, 0))
            self.assertEqual(
                (picture.width, picture.height),
                (exported.slide_width, exported.slide_height),
            )
        self.assertTrue(render_mock.call_args.kwargs["include_screenshots"])
        self.assertFalse(render_mock.call_args.kwargs["include_stylesheet"])
        compile_mock.assert_not_called()
        writer_mock.assert_not_called()
        watermark_mock.assert_called_once()
        self.assertEqual(COMPILER_CONTRACT, "htmlir-semantic-schema-v4-v6")
        self.assertEqual(WRITER_CONTRACT, "pptx2html-source-free-schema-native-fonts-v6")

    def test_default_pipeline_is_editable_and_embeds_deck_fonts(self) -> None:
        deck = {
            "version": 4,
            "deck": {"title": "test", "slides": [{"id": "s1"}]},
        }
        rendered = {
            "success": True,
            "renderer_contract": RENDERER_CONTRACT,
            "stylesheet": ".schema-slide{width:1280px;height:720px}",
            "slides": [
                {
                    "root_html": '<section class="schema-slide"></section>',
                    "native_components": [{"id": "c1", "type": "table"}],
                }
            ],
        }
        ir = {"name": "compiled", "slides": [{"name": "s1", "shapes": []}]}
        with (
            patch("schema_v4_export._render_schema_deck", return_value=rendered),
            patch("schema_v4_export.parse_html_to_ir", return_value={"presentation": ir}) as compile_mock,
            patch("schema_v4_export.export_ir_to_pptx", return_value=b"legacy-writer") as writer_mock,
            patch(
                "schema_v4_export.apply_schema_v4_native_components",
                return_value=b"native-writer",
            ) as native_mock,
            patch(
                "schema_v4_export.embed_schema_v4_fonts",
                return_value=b"font-writer",
            ) as font_mock,
        ):
            output = export_schema_v4_deck_to_pptx(deck, export_watermark=True)
        self.assertEqual(output, b"font-writer")
        self.assertTrue(compile_mock.call_args.kwargs["render_mode"] == "semantic")
        self.assertEqual(
            compile_mock.call_args.kwargs["root_selector"],
            ".schema-slide",
        )
        self.assertTrue(writer_mock.call_args.kwargs["export_watermark"])
        self.assertEqual(native_mock.call_args.args[1], [[{"id": "c1", "type": "table"}]])
        self.assertEqual(font_mock.call_args.args[0], b"native-writer")
        self.assertEqual(RENDERER_CONTRACT, "yrkie-doe-schema-v4-server-renderer-v101")
        self.assertEqual((RUNTIME_ROOT / "schema_v4_renderer.js").read_text(encoding="utf-8").count(RENDERER_CONTRACT), 3)
        self.assertEqual(COMPILER_CONTRACT, "htmlir-semantic-schema-v4-v6")
        self.assertEqual(
            WRITER_CONTRACT,
            "pptx2html-source-free-schema-native-fonts-v6",
        )

    def test_workspace_thumbnail_renders_only_the_first_slide_at_double_density(self) -> None:
        deck = {
            "version": 4,
            "deck": {"title": "cover", "slides": [{"id": "s1"}]},
        }
        rendered = {
            "success": True,
            "renderer_contract": RENDERER_CONTRACT,
            "slides": [{"screenshot_png_base64": self._screenshot_base64()}],
        }
        with patch(
            "schema_v4_export._render_schema_deck",
            return_value=rendered,
        ) as render_mock:
            thumbnail = render_schema_v4_thumbnail(deck, total_slides=12)

        with Image.open(BytesIO(thumbnail)) as image:
            self.assertEqual(image.format, "WEBP")
            self.assertEqual(image.size, (640, 360))
        self.assertEqual(render_mock.call_args.kwargs["slide_indices"], [0])
        self.assertEqual(render_mock.call_args.kwargs["total_slides"], 12)
        self.assertTrue(render_mock.call_args.kwargs["include_screenshots"])
        self.assertFalse(render_mock.call_args.kwargs["include_stylesheet"])

    def test_three_line_data_table_rebound_uses_border_box_height(self) -> None:
        deck = {
            "version": 4,
            "pageNumber": {"enabled": True, "format": "fraction"},
            "deck": {
                "title": "Table fit regression",
                "theme": "yrkie-theme",
                "slides": [
                    {
                        "id": "slide-08",
                        "size": {"width": 960, "height": 540, "unit": "pt"},
                        "theme": "business",
                        "editor": True,
                        "layout": {
                            "preset": "flow",
                            "spacing": {
                                "groupGap": "18pt",
                                "sectionGap": "18pt",
                                "internalGap": "12pt",
                                "pagePadding": "36pt",
                            },
                            "fallback": "needs-continuation",
                            "autoRepair": True,
                            "allowsBleed": False,
                            "bleedRegion": None,
                            "lockSafeArea": False,
                            "overflowPolicy": "reject-edit",
                            "maxRepairPasses": 4,
                            "allowEmergency8Pt": False,
                        },
                        "background": None,
                        "blocks": [
                            {
                                "id": "s8-title",
                                "type": "heading",
                                "props": {
                                    "text": (
                                        "Adapting the organizational structure to "
                                        "regional agility and responsiveness"
                                    ),
                                    "level": 1,
                                    "kicker": "OPERATING MODEL EVOLUTION",
                                },
                                "layout": {
                                    "lock": None,
                                    "span": "auto",
                                    "region": None,
                                    "textFit": {
                                        "mode": "readable",
                                        "minBodyPt": 11,
                                        "optionalMinBodyPt": 11,
                                    },
                                    "priority": 50,
                                    "slotSpace": {
                                        "top": "6pt",
                                        "left": 0,
                                        "right": 0,
                                        "bottom": "6pt",
                                    },
                                },
                            },
                            {
                                "id": "s8-table",
                                "type": "table",
                                "props": {
                                    "content": {
                                        "columns": [
                                            {
                                                "header": "Operating Dimension",
                                                "align": "left",
                                            },
                                            {
                                                "header": "Historical Direct Model",
                                                "align": "left",
                                            },
                                            {
                                                "header": "Evolved Agile Model",
                                                "align": "left",
                                            },
                                        ],
                                        "rows": [
                                            {
                                                "cells": [
                                                    "Decision Architecture",
                                                    (
                                                        "Centralized headquarters command "
                                                        "structure"
                                                    ),
                                                    (
                                                        "Empowered regional managers with "
                                                        "real-time flexibility"
                                                    ),
                                                ]
                                            },
                                            {
                                                "cells": [
                                                    "Menu & Merchandising",
                                                    "Uniform nationwide menu rollout cadence",
                                                    (
                                                        "Hyper-localized seasonal items and "
                                                        "regional pricing tests"
                                                    ),
                                                ]
                                            },
                                            {
                                                "cells": [
                                                    "Supply Integration",
                                                    "Global standardized sourcing cycles",
                                                    (
                                                        "Direct linkage to CIP and domestic "
                                                        "ingredient partners"
                                                    ),
                                                ]
                                            },
                                            {
                                                "cells": [
                                                    "Frontline Authority",
                                                    (
                                                        "Rigid operational playbooks and "
                                                        "guidelines"
                                                    ),
                                                    (
                                                        "Dynamic scheduling and micro-market "
                                                        "promotional empowerment"
                                                    ),
                                                ]
                                            },
                                        ],
                                    },
                                    "borderStyle": "three-line",
                                    "columnSizing": "auto",
                                },
                                "layout": {
                                    "lock": None,
                                    "span": "auto",
                                    "region": None,
                                    "textFit": {
                                        "mode": "readable",
                                        "minBodyPt": 11,
                                        "optionalMinBodyPt": 11,
                                    },
                                    "priority": 50,
                                    "slotSpace": {
                                        "top": "6pt",
                                        "left": 0,
                                        "right": 0,
                                        "bottom": "6pt",
                                    },
                                },
                            },
                        ],
                    }
                ],
            },
        }

        rendered = _render_schema_deck(
            deck,
            {},
            timeout_seconds=30,
            include_screenshots=False,
            include_stylesheet=False,
        )
        root_html = rendered["slides"][0]["root_html"]

        self.assertIn('data-lines="three-line"', root_html)
        self.assertIn('data-overflow="false"', root_html)
        self.assertNotIn('data-overflow="true"', root_html)
        self.assertIn(
            "Dynamic scheduling and micro-market promotional empowerment",
            root_html,
        )

    def test_mathml_snapshot_compiles_to_editable_omml(self) -> None:
        component = {
            "id": "formula-1",
            "type": "math-formula",
            "props": {
                "content": {
                    "formula": r"\frac{x^2}{n}",
                    "align": "center",
                }
            },
            "bounds_pt": {"left": 10, "top": 20, "width": 300, "height": 80},
            "mathml": (
                '<math><mfrac><msup><mi>x</mi><mn>2</mn></msup>'
                '<mi>n</mi></mfrac></math>'
            ),
            "font_size_pt": 24,
            "color": "rgb(21, 45, 79)",
        }
        spec = formula_spec(component)
        self.assertEqual(spec.color, "152D4F")
        omath = _mathml_to_omml(spec.mathml)
        self.assertEqual(omath.tag, f"{{{MATH_NS}}}oMath")
        self.assertEqual(len(omath.findall(f".//{{{MATH_NS}}}f")), 1)
        self.assertEqual(len(omath.findall(f".//{{{MATH_NS}}}sSup")), 1)

        overline = _mathml_to_omml(
            '<math><mrow class="ppt-math-formula__overline"><mi>x</mi></mrow></math>'
        )
        self.assertEqual(len(overline.findall(f".//{{{MATH_NS}}}bar")), 1)

    def test_overline_stays_editable_without_office_xslt(self) -> None:
        for mathml in (
            '<math><mrow class="ppt-math-formula__overline"><mi>x</mi></mrow></math>',
            '<math><mover accent="true"><mi>x</mi><mo>‾</mo></mover></math>',
        ):
            with self.subTest(mathml=mathml), patch("schema_v4_math._office_xslt_path", return_value=None):
                omath = _mathml_to_omml(mathml)
            self.assertEqual(len(omath.findall(f".//{{{MATH_NS}}}bar")), 1)
            self.assertEqual(omath.find(f".//{{{MATH_NS}}}barPr/{{{MATH_NS}}}pos").get(f"{{{MATH_NS}}}val"), "top")
            self.assertEqual([node.text for node in omath.findall(f".//{{{MATH_NS}}}t")], ["x"])

    def test_editable_chart_legends_stay_below_the_plot_and_use_the_combined_bounds(self) -> None:
        chart_contents = {
            "grouped-bar-chart": {
                "categories": ["2024", "2025"],
                "series": [{"name": "Core", "data": [48, 75], "color": "#203754"}],
            },
            "line-chart": {
                "categories": ["FY1", "FY2"],
                "series": [{"name": "Base", "data": [20, 45], "color": "#203754"}],
            },
            "stacked-bar-chart": {
                "categories": ["2024", "2025"],
                "series": [{"name": "Core", "data": [48, 75], "color": "#203754"}],
            },
            "radar-chart": {
                "axes": [{"label": "Product"}, {"label": "Delivery"}],
                "datasets": [{"name": "Current", "values": [82, 88], "color": "#203754"}],
            },
            "donut-chart": {
                "items": [
                    {"label": "Core", "value": 62, "color": "#203754"},
                    {"label": "Growth", "value": 38, "color": "#bd2d24"},
                ]
            },
        }
        for component_type, content in chart_contents.items():
            with self.subTest(component_type=component_type):
                presentation = Presentation()
                presentation.slides.add_slide(presentation.slide_layouts[6])
                source = BytesIO()
                presentation.save(source)
                component = {
                    "id": f"{component_type}-1",
                    "type": component_type,
                    "props": {
                        "title": "Legacy component title",
                        "showLegend": True,
                        "content": content,
                    },
                    "preserve_component_text": True,
                    "bounds_pt": {"left": 100, "top": 100, "width": 300, "height": 200},
                    "legend_bounds_pt": {
                        "left": 100,
                        "top": 310,
                        "width": 300,
                        "height": 30,
                    },
                }

                output = apply_schema_v4_native_components(source.getvalue(), [[component]])
                exported = Presentation(BytesIO(output))
                chart_shape = next(shape for shape in exported.slides[0].shapes if shape.has_chart)

                self.assertFalse(chart_shape.chart.has_title)
                self.assertEqual(chart_shape.chart.legend.position, XL_LEGEND_POSITION.BOTTOM)
                self.assertFalse(chart_shape.chart.legend.include_in_layout)
                self.assertAlmostEqual(chart_shape.left / 12_700, 100, places=2)
                self.assertAlmostEqual(chart_shape.top / 12_700, 100, places=2)
                self.assertAlmostEqual(chart_shape.width / 12_700, 300, places=2)
                self.assertAlmostEqual(chart_shape.height / 12_700, 240, places=2)

    def test_reversed_stacked_legend_preserves_the_bottom_html_fallback(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        legend_fallback = slide.shapes.add_textbox(Pt(100), Pt(310), Pt(300), Pt(30))
        legend_fallback.text = "Other · New · Core"
        source = BytesIO()
        presentation.save(source)
        component = {
            "id": "stacked-reversed",
            "type": "stacked-bar-chart",
            "props": {
                "showLegend": True,
                "content": {
                    "categories": ["2024", "2025"],
                    "legendReversed": True,
                    "series": [
                        {"name": "Core", "data": [48, 75], "color": "#203754"},
                        {"name": "New", "data": [18, 38], "color": "#bd2d24"},
                        {"name": "Other", "data": [10, 12], "color": "#d8dadd"},
                    ],
                },
            },
            "bounds_pt": {"left": 100, "top": 100, "width": 300, "height": 200},
            "legend_bounds_pt": {"left": 100, "top": 310, "width": 300, "height": 30},
        }

        output = apply_schema_v4_native_components(source.getvalue(), [[component]])
        exported = Presentation(BytesIO(output))
        shapes = list(exported.slides[0].shapes)
        chart_shape = next(shape for shape in shapes if shape.has_chart)

        self.assertFalse(chart_shape.chart.has_legend)
        self.assertAlmostEqual(chart_shape.height / 12_700, 200, places=2)
        self.assertIn(
            "Other · New · Core",
            [shape.text for shape in shapes if getattr(shape, "has_text_frame", False)],
        )

    def test_font_subset_scope_uses_visible_text_and_actual_faces(self) -> None:
        presentation = {
            "slides": [
                {
                    "shapes": [
                        {
                            "text_frame": {
                                "paragraphs": [
                                    {
                                        "runs": [
                                            {
                                                "text": "正文甲",
                                                "font": {
                                                    "family": "Infaxiswield Humanist",
                                                    "bold": False,
                                                    "italic": False,
                                                },
                                            },
                                            {
                                                "text": "标题乙",
                                                "font": {
                                                    "family": "Infaxiswield Humanist",
                                                    "bold": True,
                                                    "italic": False,
                                                },
                                            },
                                        ]
                                    }
                                ]
                            }
                        }
                    ]
                }
            ]
        }
        characters, faces = _collect_visible_characters(
            presentation,
            native_components_by_slide=[[]],
        )
        self.assertTrue(set("正文甲标题乙").issubset(set(characters)))
        self.assertEqual(
            faces,
            {
                ("infaxiswield humanist", "regular"),
                ("infaxiswield humanist", "bold"),
            },
        )

    def test_schema_font_usage_separates_misans_and_alexandria(self) -> None:
        presentation = {
            "slides": [
                {
                    "shapes": [
                        {
                            "text": {
                                "paragraphs": [
                                    {
                                        "runs": [
                                            {
                                                "text": "中文正文",
                                                "font": {"family": "MiSans"},
                                            },
                                            {
                                                "text": "English Heading",
                                                "font": {
                                                    "family": "Alexandria",
                                                    "bold": True,
                                                },
                                            },
                                        ]
                                    }
                                ]
                            }
                        }
                    ]
                }
            ]
        }

        characters, faces = _collect_visible_characters(
            presentation,
            native_components_by_slide=[[]],
        )

        self.assertTrue(set("中文正文English Heading").issubset(set(characters)))
        self.assertEqual(
            faces,
            {("misans", "regular"), ("alexandria", "bold")},
        )

    def test_schema_font_usage_maps_synthetic_italic_to_available_family_face(self) -> None:
        presentation = {
            "slides": [
                {
                    "shapes": [
                        {
                            "text": {
                                "paragraphs": [
                                    {
                                        "runs": [
                                            {
                                                "text": "中文斜体",
                                                "font": {
                                                    "family": "MiSans",
                                                    "italic": True,
                                                },
                                            },
                                            {
                                                "text": "English Bold Italic",
                                                "font": {
                                                    "family": "Alexandria",
                                                    "bold": True,
                                                    "italic": True,
                                                },
                                            },
                                        ]
                                    }
                                ]
                            }
                        }
                    ]
                }
            ]
        }

        _, faces = _collect_visible_characters(
            presentation,
            native_components_by_slide=[[]],
        )

        self.assertEqual(faces, {("misans", "regular"), ("alexandria", "bold")})

    def test_schema_font_usage_keeps_medium_weight_as_distinct_regular_families(self) -> None:
        presentation = {
            "slides": [
                {
                    "shapes": [
                        {
                            "text": {
                                "paragraphs": [
                                    {
                                        "runs": [
                                            {
                                                "text": "中文半粗",
                                                "font": {
                                                    "family": "MiSans Medium",
                                                    "weight": 400,
                                                },
                                            },
                                            {
                                                "text": "English Medium",
                                                "font": {
                                                    "family": "Alexandria Medium",
                                                    "weight": 400,
                                                },
                                            },
                                        ]
                                    }
                                ]
                            }
                        }
                    ]
                }
            ]
        }

        _, faces = _collect_visible_characters(
            presentation,
            native_components_by_slide=[[]],
        )

        self.assertEqual(
            faces,
            {("misans medium", "regular"), ("alexandria medium", "regular")},
        )

    def test_schema_font_faces_use_the_300_400_600_contract(self) -> None:
        renderer_root = (
            DOE_PUBLIC_ROOT
        )
        faces = {
            (face.family, face.slot): (face.weight, face.path.name)
            for face in trusted_schema_v4_fonts(renderer_root)
            if face.family.startswith("MiSans") or face.family.startswith("Alexandria")
        }
        self.assertEqual(faces[("MiSans", "regular")], (300, "MiSans-Light.woff2"))
        self.assertEqual(faces[("MiSans Medium", "regular")], (400, "MiSans-Medium.woff2"))
        self.assertEqual(faces[("MiSans", "bold")], (600, "MiSans-Semibold.woff2"))
        self.assertEqual(faces[("Alexandria", "regular")][0], 300)
        self.assertEqual(faces[("Alexandria Medium", "regular")][0], 400)
        self.assertEqual(faces[("Alexandria", "bold")][0], 600)

    def test_native_charts_choose_font_from_visible_script(self) -> None:
        chinese = {"type": "grouped-bar-chart", "props": {"title": "季度收入"}}
        english = {"type": "grouped-bar-chart", "props": {"title": "Quarterly Revenue"}}
        self.assertEqual(_component_font_name(chinese), "MiSans")
        self.assertEqual(_component_font_name(english), "Alexandria")
        _, chinese_faces = _collect_visible_characters(
            {"slides": []},
            native_components_by_slide=[[chinese]],
        )
        _, english_faces = _collect_visible_characters(
            {"slides": []},
            native_components_by_slide=[[english]],
        )
        self.assertEqual(
            chinese_faces,
            {("misans", "regular"), ("misans medium", "regular")},
        )
        self.assertEqual(
            english_faces,
            {("alexandria", "regular"), ("alexandria medium", "regular")},
        )

    def test_font_parts_are_added_to_the_existing_pptx_package(self) -> None:
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        source = BytesIO()
        presentation.save(source)
        output = _embed_font_parts(
            source.getvalue(),
            [
                PreparedFontFace(
                    family="Infaxiswield Humanist",
                    style="normal",
                    weight=400,
                    slot="regular",
                    font_part_bytes=b"font-part",
                    panose="",
                    used_codepoints=1,
                    missing_codepoints=(),
                )
            ],
        )
        with ZipFile(BytesIO(output), "r") as package:
            self.assertEqual(package.read("ppt/fonts/font1.fntdata"), b"font-part")
            self.assertIn(b"embeddedFontLst", package.read("ppt/presentation.xml"))
            self.assertIn(
                b"/relationships/font",
                package.read("ppt/_rels/presentation.xml.rels"),
            )


if __name__ == "__main__":
    unittest.main()

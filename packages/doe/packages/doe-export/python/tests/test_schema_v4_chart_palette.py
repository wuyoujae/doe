from __future__ import annotations

from test_paths import DOE_PUBLIC_ROOT

import sys
import unittest
from io import BytesIO
from pathlib import Path

from pptx import Presentation


RUNTIME_ROOT = Path(__file__).resolve().parents[1]
if str(RUNTIME_ROOT) not in sys.path:
    sys.path.insert(0, str(RUNTIME_ROOT))

from schema_v4_native import apply_schema_v4_native_components  # noqa: E402
from schema_v4_math import formula_spec  # noqa: E402


class SchemaV4ChartPaletteTests(unittest.TestCase):
    @staticmethod
    def _native_chart(component_type: str, series: list[dict[str, object]]):
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        source = BytesIO()
        presentation.save(source)
        component = {
            "id": f"{component_type}-1",
            "type": component_type,
            "props": {
                "title": "Chart title",
                "showLegend": True,
                "content": {
                    "categories": ["2024", "2025"],
                    "series": series,
                },
            },
            "bounds_pt": {"left": 100, "top": 100, "width": 360, "height": 220},
        }
        output = apply_schema_v4_native_components(source.getvalue(), [[component]])
        exported = Presentation(BytesIO(output))
        return next(shape.chart for shape in exported.slides[0].shapes if shape.has_chart)

    def test_grouped_and_line_charts_share_blue_defaults_and_preserve_custom_colors(self) -> None:
        series = [
            {"name": "Primary", "data": [48, 75]},
            {"name": "Secondary", "color": "  ", "data": [18, 38]},
            {"name": "Custom", "color": "#123456", "data": [12, 16]},
        ]
        for component_type in ("grouped-bar-chart", "line-chart"):
            with self.subTest(component_type=component_type):
                chart = self._native_chart(component_type, series)
                self.assertEqual(
                    [str(item.format.line.color.rgb) for item in chart.series],
                    ["1B54DA", "D2DDF8", "123456"],
                )
                if component_type == "grouped-bar-chart":
                    self.assertEqual(
                        [str(item.format.fill.fore_color.rgb) for item in chart.series],
                        ["1B54DA", "D2DDF8", "123456"],
                    )
                else:
                    self.assertEqual(
                        [str(item.marker.format.fill.fore_color.rgb) for item in chart.series],
                        ["1B54DA", "D2DDF8", "123456"],
                    )

    def test_native_chart_text_uses_heading_and_body_defaults(self) -> None:
        for component_type in ("grouped-bar-chart", "line-chart"):
            with self.subTest(component_type=component_type):
                chart = self._native_chart(
                    component_type,
                    [{"name": "Primary", "data": [48, 75]}],
                )
                title_run = chart.chart_title.text_frame.paragraphs[0].runs[0]
                self.assertEqual(str(title_run.font.color.rgb), "1B1B27")
                self.assertEqual(str(chart.legend.font.color.rgb), "404155")
                self.assertEqual(
                    str(chart.value_axis.tick_labels.font.color.rgb),
                    "404155",
                )
                self.assertEqual(
                    str(chart.category_axis.tick_labels.font.color.rgb),
                    "404155",
                )
                if component_type == "grouped-bar-chart":
                    self.assertEqual(
                        str(chart.plots[0].data_labels.font.color.rgb),
                        "404155",
                    )

    def test_stacked_chart_uses_the_extended_blue_defaults(self) -> None:
        chart = self._native_chart(
            "stacked-bar-chart",
            [
                {"name": "Primary", "data": [48, 75]},
                {"name": "Light", "data": [18, 38]},
                {"name": "Deep", "data": [12, 16]},
            ],
        )
        self.assertEqual(
            [str(item.format.line.color.rgb) for item in chart.series],
            ["1B54DA", "D2DDF8", "10317F"],
        )

    def test_formula_uses_heading_default_and_preserves_authored_color(self) -> None:
        component = {
            "id": "formula-1",
            "type": "math-formula",
            "props": {"content": {"formula": "x^2"}},
            "bounds_pt": {"left": 10, "top": 20, "width": 300, "height": 80},
            "mathml": "<math><msup><mi>x</mi><mn>2</mn></msup></math>",
        }
        self.assertEqual(formula_spec(component).color, "1B1B27")
        component["props"]["content"]["color"] = "#123456"
        self.assertEqual(formula_spec(component).color, "123456")


if __name__ == "__main__":
    unittest.main()

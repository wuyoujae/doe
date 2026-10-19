"""The shared renderer preserves complete component scaling in native exports."""
from io import BytesIO
import base64
import re
import shutil
import unittest
import fitz
from pptx import Presentation
from schema_v4_export import export_schema_v4_deck_to_pptx, _render_schema_deck
from test_paths import DOE_PUBLIC_ROOT


def compression_deck():
    return {'version': 4, 'deck': {'theme': 'yrkie-theme', 'slides': [{'id': 'scale', 'blocks': [
        {'id': 'group', 'type': 'layout-group', 'layout': {'variant': 'row', 'columns': 3}, 'children': [
            {'id': 'cards', 'type': 'contents', 'props': {'presentation': 'cards:leaf', 'items': [
                {'title': f'Card {i}', 'text': f'Complete content {i}'} for i in range(1, 11)]}},
            {'id': 'peer', 'type': 'text', 'props': {'text': 'Unaffected peer'}},
            {'id': 'peer2', 'type': 'text', 'props': {'text': 'Another peer'}},
        ]},
    ]}]}}


@unittest.skipUnless(shutil.which('node') and DOE_PUBLIC_ROOT.is_dir(), 'built DOE runtime required')
class FlowCompressionExportTests(unittest.TestCase):
    def test_fixed_radial_geometry_and_text_share_the_export_scale(self):
        shape = {'id': 'fixed', 'type': 'contents', 'props': {'presentation': 'shape:semicircle', 'items': [
            {'title': f'Item {i}', 'text': f'Original content {i}'} for i in range(1, 4)]}}
        deck = {'version': 4, 'deck': {'theme': 'yrkie-theme', 'slides': [{'id': 'fixed-export', 'blocks': [
            {'id': 'group', 'type': 'layout-group', 'layout': {'variant': 'row', 'columns': 2},
             'children': [shape, {'id': 'peer', 'type': 'text', 'props': {'text': 'Unaffected peer'}}]},
        ]}]}}
        rendered = _render_schema_deck(deck, {}, timeout_seconds=120, include_screenshots=False, include_stylesheet=False)
        root = re.search(r'<[^>]*data-component-id="fixed"[^>]*>', rendered['slides'][0]['root_html']).group(0)
        ratio = float(re.search(r'data-flow-paint-scale="([^"]+)"', root).group(1))
        self.assertLess(ratio, 1)
        self.assertAlmostEqual(float(re.search(r'width:\s*([\d.]+)px', root).group(1)), 960, delta=.05)
        output = export_schema_v4_deck_to_pptx(deck, embed_fonts=False)
        presentation = Presentation(BytesIO(output))
        runs = [run for item in presentation.slides[0].shapes if item.has_text_frame
                for paragraph in item.text_frame.paragraphs for run in paragraph.runs]
        combined = re.sub(r'\s+', '', ''.join(run.text for run in runs))
        for i in range(1, 4):
            self.assertIn(f'Originalcontent{i}', combined)
        content_runs = [run for run in runs if 'Original' in run.text]
        self.assertEqual(len(content_runs), 3)
        for run in content_runs:
            self.assertAlmostEqual(run.font.size.pt, 15.5 * ratio, delta=.08)
        peer = next(run for run in runs if 'Unaffected peer' in run.text)
        self.assertAlmostEqual(peer.font.size.pt, 15.5, delta=.08)

    def test_whole_component_export_preserves_all_text_and_sub_8pt_fonts(self):
        deck = compression_deck()
        rendered = _render_schema_deck(deck, {}, timeout_seconds=120, include_screenshots=False, include_stylesheet=False)
        html = rendered['slides'][0]['root_html']
        root = re.search(r'<[^>]*data-component-id="cards"[^>]*>', html).group(0)
        ratio = float(re.search(r'data-flow-paint-scale="([^"]+)"', root).group(1))
        self.assertGreaterEqual(ratio, 1 / 3)
        self.assertLess(ratio, .52)
        output = export_schema_v4_deck_to_pptx(deck, embed_fonts=False)
        presentation = Presentation(BytesIO(output))
        runs = [run for shape in presentation.slides[0].shapes if shape.has_text_frame
                for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
        text = ''.join(run.text for run in runs)
        for i in range(1, 11):
            self.assertIn(f'Complete content {i}', text)
        card_runs = [run for run in runs if 'Complete content' in run.text]
        self.assertTrue(card_runs)
        for run in card_runs:
            self.assertAlmostEqual(run.font.size.pt, 15.5 * ratio, delta=.08)
        peer = next(run for run in runs if 'Unaffected peer' in run.text)
        self.assertAlmostEqual(peer.font.size.pt, 15.5, delta=.08)
        pdf = _render_schema_deck(deck, {}, timeout_seconds=120, include_screenshots=False,
                                  include_stylesheet=False, output_format='pdf')
        with fitz.open(stream=base64.b64decode(pdf['pdf_base64']), filetype='pdf') as document:
            self.assertEqual(len(document), 1)
            text = document[0].get_text()
            for i in range(1, 11):
                self.assertIn(f'Complete content {i}', text)

    def test_internal_content_exports_ellipsis_and_preserves_source(self):
        deck = {'version': 4, 'deck': {'theme': 'yrkie-theme', 'slides': [{'id': 'truncated', 'blocks': [
            {'id': 'text', 'type': 'text', 'props': {'text': 'Complete content. ' * 3000}}]}]}}
        output = export_schema_v4_deck_to_pptx(deck, embed_fonts=False)
        presentation = Presentation(BytesIO(output))
        visible = ''.join(shape.text for shape in presentation.slides[0].shapes if shape.has_text_frame)
        self.assertIn('...', visible)
        self.assertLess(len(visible), len(deck['deck']['slides'][0]['blocks'][0]['props']['text']))
        self.assertEqual(deck['deck']['slides'][0]['blocks'][0]['props']['text'], 'Complete content. ' * 3000)

if __name__ == '__main__':
    unittest.main()

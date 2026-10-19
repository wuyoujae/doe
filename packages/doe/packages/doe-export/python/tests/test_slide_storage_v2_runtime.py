from __future__ import annotations

import base64
import dataclasses
import hashlib
import http.client
import io
import json
import re
import struct
import sys
import threading
import unittest
from pathlib import Path
from urllib.parse import quote

from PIL import Image as PILImage
from pptx import Presentation as PptxPresentation


RUNTIME_ROOT = Path(__file__).resolve().parents[1]
if str(RUNTIME_ROOT) not in sys.path:
    sys.path.insert(0, str(RUNTIME_ROOT))

from htmlir_server import HtmlIrHttpServer, _allowed_hosts  # noqa: E402
from pptx_pipeline import (  # noqa: E402
    extract_editable_slide_html,
    import_pptx_to_canonical,
    render_ir_to_canonical,
    render_ir_to_slide_document,
)
from pptx2html.asset_content import css_font_format  # noqa: E402
from pptx2html.ir.builder import PresentationBuilder  # noqa: E402
from pptx2html.ir.model import (  # noqa: E402
    Bounds,
    Border,
    EmbeddedFont,
    Fill,
    ImageInfo,
    Paragraph,
    Presentation,
    Run,
    Shape,
    Slide,
    TableCell,
    TableInfo,
    TextFrame,
)


def _presentation_with_inline_assets() -> Presentation:
    image = b"\x89PNG\r\n\x1a\nv2-image"
    font = b"\x00\x01\x00\x00v2-font"
    return Presentation(
        width=960.0,
        height=540.0,
        embedded_fonts=[
            EmbeddedFont(
                typeface="V2 Test",
                content_type="font/otf",
                data=base64.b64encode(font).decode("ascii"),
            )
        ],
        slides=[
            Slide(
                index=0,
                name="V2 slide",
                shapes=[
                    Shape(
                        type="picture",
                        name="Stable picture",
                        shape_id=17,
                        bounds=Bounds(12.0, 18.0, 320.0, 180.0),
                        image=ImageInfo(
                            content_type="image/png",
                            data=base64.b64encode(image).decode("ascii"),
                            alpha=0.65,
                        ),
                    )
                ],
            )
        ],
    )


def _png(color: tuple[int, int, int], size: tuple[int, int] = (4, 2)) -> bytes:
    output = io.BytesIO()
    PILImage.new("RGB", size, color).save(output, format="PNG")
    return output.getvalue()


def _text_frame(value: str) -> TextFrame:
    return TextFrame(paragraphs=[Paragraph(runs=[Run(text=value)])])


class SlideStorageV2RuntimeTests(unittest.TestCase):
    def test_render_returns_pure_html_and_externalized_current_ir(self) -> None:
        result = render_ir_to_slide_document(
            dataclasses.asdict(_presentation_with_inline_assets())
        )

        self.assertNotIn("canonical_html", result)
        self.assertTrue(result["editable_html"].startswith("<section "))
        self.assertNotIn("<html", result["editable_html"].lower())
        self.assertNotIn("<script", result["editable_html"].lower())
        self.assertIn('data-yrkie-element-id="pptx-17"', result["editable_html"])
        self.assertIn("yrkie-asset-data://", result["editable_html"])
        self.assertNotIn("data:image/", result["editable_html"])
        self.assertEqual(len(result["assets"]), 2)
        self.assertEqual(
            result["current_ir"]["embedded_fonts"][0]["content_type"],
            "font/ttf",
        )
        serialized_ir = json.dumps(result["current_ir"])
        self.assertNotIn(";base64,", serialized_ir)
        self.assertIn("yrkie-asset-base64://", serialized_ir)
        self.assertEqual(
            result["editable_html_sha256"],
            hashlib.sha256(result["editable_html"].encode("utf-8")).hexdigest(),
        )
        storage_bytes = json.dumps(
            result["current_ir"],
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
        self.assertEqual(
            result["current_ir_sha256"],
            hashlib.sha256(storage_bytes).hexdigest(),
        )

    def test_render_preserves_data_url_shaped_visible_text(self) -> None:
        literal = "data:image/png;base64,AA== is documentation, not an asset"
        presentation = Presentation(
            width=960,
            height=540,
            slides=[
                Slide(
                    index=0,
                    shapes=[
                        Shape(
                            type="auto_shape",
                            shape_id=9,
                            bounds=Bounds(10, 10, 600, 60),
                            text=TextFrame(
                                paragraphs=[
                                    Paragraph(runs=[Run(text=literal)])
                                ]
                            ),
                        )
                    ],
                )
            ],
        )

        result = render_ir_to_slide_document(dataclasses.asdict(presentation))

        self.assertIn(literal, result["editable_html"])
        self.assertEqual(
            result["current_ir"]["slides"][0]["shapes"][0]["text"]
            ["paragraphs"][0]["runs"][0]["text"],
            literal,
        )

    def test_canonical_html_is_only_returned_on_demand(self) -> None:
        result = render_ir_to_slide_document(
            dataclasses.asdict(_presentation_with_inline_assets()),
            include_canonical_html=True,
        )

        self.assertIn('id="pptx2html-canonical-ir"', result["canonical_html"])
        self.assertEqual(
            result["canonical_html_sha256"],
            hashlib.sha256(result["canonical_html"].encode("utf-8")).hexdigest(),
        )

    def test_extract_does_not_return_stale_canonical_ir(self) -> None:
        rendered = render_ir_to_canonical(
            dataclasses.asdict(_presentation_with_inline_assets())
        )
        modified = rendered["html"].replace(
            'data-slide-name="V2 slide"',
            'data-slide-name="AI edited slide"',
            1,
        )

        extracted = extract_editable_slide_html(modified)

        self.assertFalse(extracted["canonical_ir_valid"])
        self.assertNotIn("current_ir", extracted)
        self.assertNotIn("current_ir_sha256", extracted)
        self.assertIn("AI edited slide", extracted["editable_html"])
        self.assertNotIn("pptx2html-canonical-ir", extracted["editable_html"])

    def test_extract_externalizes_legacy_visible_data_url(self) -> None:
        raw = b"legacy-image"
        encoded = base64.b64encode(raw).decode("ascii")
        canonical = f"""
<!doctype html><html><body>
<!-- PPTX2HTML_SLIDE_START:0 -->
<section data-pptx2html-slide="1"><img src="data:image/png;base64,{encoded}"></section>
<!-- PPTX2HTML_SLIDE_END:0 -->
</body></html>
"""

        extracted = extract_editable_slide_html(canonical)

        digest = hashlib.sha256(raw).hexdigest()
        self.assertFalse(extracted["canonical_ir_valid"])
        self.assertIn(f"yrkie-asset-data://{digest}", extracted["editable_html"])
        self.assertEqual(extracted["assets"][0]["asset_id"], digest)

    def test_extract_rejects_unexternalizable_inline_data_url(self) -> None:
        canonical = """
<!-- PPTX2HTML_SLIDE_START:0 -->
<section data-pptx2html-slide="1"><img src="data:image/png;base64,***"></section>
<!-- PPTX2HTML_SLIDE_END:0 -->
"""
        with self.assertRaisesRegex(ValueError, "inline data URLs"):
            extract_editable_slide_html(canonical)

    def test_composed_picture_svg_is_an_external_asset_not_html_token_bloat(self) -> None:
        image = b"\x89PNG\r\n\x1a\ncomposed-image"
        presentation = Presentation(
            width=960,
            height=540,
            slides=[
                Slide(
                    index=0,
                    shapes=[
                        Shape(
                            type="auto_shape",
                            name="Picture fill",
                            shape_id=9,
                            bounds=Bounds(0, 0, 320, 180),
                            fill=Fill(
                                type="picture",
                                image=ImageInfo(
                                    content_type="image/png",
                                    data=base64.b64encode(image).decode("ascii"),
                                    alpha=0.65,
                                ),
                            ),
                        )
                    ],
                )
            ],
        )
        result = render_ir_to_slide_document(
            dataclasses.asdict(presentation)
        )

        editable = result["editable_html"]
        self.assertNotIn("data:image/svg+xml", editable)
        self.assertNotIn("%3Csvg", editable)
        self.assertLess(len(editable), 12_000)
        svg_assets = [
            item for item in result["assets"]
            if item["content_type"] == "image/svg+xml"
        ]
        self.assertEqual(len(svg_assets), 1)
        svg = base64.b64decode(svg_assets[0]["data_base64"]).decode("utf-8")
        self.assertIn("data:image/png;base64,", svg)
        self.assertNotIn("yrkie-asset-", svg)

    def test_large_percent_encoded_svg_is_externalized_from_editable_html(self) -> None:
        points = " ".join(f"L {index} {index % 300}" for index in range(30_000))
        svg = (
            '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 30000 300">'
            f'<path d="M 0 0 {points}"/></svg>'
        )
        canonical = f"""
<!-- PPTX2HTML_SLIDE_START:0 -->
<section data-pptx2html-slide="1" style="background-image:url('data:image/svg+xml,{quote(svg, safe="/:;,+=")}')"></section>
<!-- PPTX2HTML_SLIDE_END:0 -->
"""

        extracted = extract_editable_slide_html(canonical)

        self.assertLess(len(extracted["editable_html"]), 2_000)
        self.assertNotIn("data:image/svg+xml", extracted["editable_html"])
        self.assertEqual(extracted["assets"][0]["content_type"], "image/svg+xml")
        self.assertGreater(extracted["assets"][0]["size_bytes"], 100_000)

    def test_missing_nested_svg_asset_gets_placeholder_and_diagnostic(self) -> None:
        missing_id = "a" * 64
        svg = (
            '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1 1">'
            f'<image href="yrkie-asset-data://{missing_id}" width="1" height="1"/>'
            "</svg>"
        )
        canonical = f"""
<!-- PPTX2HTML_SLIDE_START:0 -->
<section data-pptx2html-slide="1" style="background-image:url('data:image/svg+xml,{quote(svg, safe="/:;,+=")}')"></section>
<!-- PPTX2HTML_SLIDE_END:0 -->
"""

        extracted = extract_editable_slide_html(
            canonical,
            known_asset_ids={missing_id},
        )

        self.assertNotIn("data:image/svg+xml", extracted["editable_html"])
        self.assertIn(
            "svg-nested-asset-missing",
            {item["code"] for item in extracted["resource_diagnostics"]},
        )

    def test_composed_svg_removes_active_css_and_nested_svg_data(self) -> None:
        nested = base64.b64encode(
            b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'
        ).decode("ascii")
        safe_png = (
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR4nGNg"
            "YAAAAAMAASsJTYQAAAAASUVORK5CYII="
        )
        svg = (
            '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 2 1">'
            "<style>@import url(https://evil.invalid/x.css);"
            ".x{background:url(data:text/html;base64,PGgxPmV2aWw8L2gxPg==)}</style>"
            "<script>alert(1)</script>"
            f'<image href="data:image/svg+xml;base64,{nested}" width="1" height="1"/>'
            f'<image href="data:image/png;base64,{safe_png}" x="1" width="1" height="1"/>'
            "</svg>"
        )
        canonical = f"""
<!-- PPTX2HTML_SLIDE_START:0 -->
<section data-pptx2html-slide="1" style="background-image:url('data:image/svg+xml,{quote(svg, safe="/:;,+=")}')"></section>
<!-- PPTX2HTML_SLIDE_END:0 -->
"""

        extracted = extract_editable_slide_html(canonical)
        svg_asset = next(
            item for item in extracted["assets"]
            if item["content_type"] == "image/svg+xml"
        )
        sanitized = base64.b64decode(svg_asset["data_base64"]).decode("utf-8")

        self.assertNotIn("<script", sanitized.lower())
        self.assertNotIn("@import", sanitized.lower())
        self.assertNotIn("data:image/svg+xml", sanitized.lower())
        self.assertIn("data:image/png;base64,", sanitized)
        self.assertIn(
            "svg-resource-sanitized",
            {item["code"] for item in extracted["resource_diagnostics"]},
        )

    def test_font_magic_maps_to_exact_mime_and_css_format(self) -> None:
        cases = {
            b"\x00\x01\x00\x00": ("font/ttf", "truetype"),
            b"OTTO": ("font/otf", "opentype"),
            b"ttcf": ("font/collection", "collection"),
            b"wOFF": ("font/woff", "woff"),
            b"wOF2": ("font/woff2", "woff2"),
        }
        for magic, (expected_mime, expected_format) in cases.items():
            with self.subTest(magic=magic):
                blob = magic + b"font-payload"
                self.assertEqual(
                    PresentationBuilder._browser_ready_font_blob(blob),
                    (expected_mime, blob),
                )
                self.assertEqual(css_font_format(expected_mime), expected_format)

                wrapper = bytearray(12)
                struct.pack_into("<I", wrapper, 4, len(blob))
                self.assertEqual(
                    PresentationBuilder._browser_ready_font_blob(
                        bytes(wrapper) + blob
                    ),
                    (expected_mime, blob),
                )

    def test_copyist_import_adds_compact_v2_slide_document_fields(self) -> None:
        source = PptxPresentation()
        slide = source.slides.add_slide(source.slide_layouts[6])
        slide.shapes.add_textbox(914400, 914400, 3657600, 914400).text = "V2"
        stream = io.BytesIO()
        source.save(stream)

        imported = import_pptx_to_canonical(stream.getvalue())

        document = imported["slides"][0]
        self.assertIn("editable_html", document)
        self.assertIn("editable_html_sha256", document)
        self.assertTrue(document["canonical_ir_valid"])
        self.assertIn("current_ir_sha256", document)
        self.assertNotIn("current_ir", document)
        self.assertIn('id="pptx2html-canonical-ir"', document["html"])
        self.assertNotIn("<script", document["editable_html"].lower())


class SlideStorageV2HttpTests(unittest.TestCase):
    TOKEN = "slide-storage-v2-runtime-test-token"

    def setUp(self) -> None:
        self.server = HtmlIrHttpServer(
            ("127.0.0.1", 0),
            runtime_token=self.TOKEN,
            allowed_hosts=_allowed_hosts("127.0.0.1", []),
            max_concurrency=1,
        )
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self) -> None:
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=2)

    def _post(self, path: str, payload: dict) -> tuple[int, dict]:
        body = json.dumps(payload).encode("utf-8")
        connection = http.client.HTTPConnection(
            "127.0.0.1",
            self.server.server_address[1],
            timeout=30,
        )
        connection.request(
            "POST",
            path,
            body=body,
            headers={
                "Authorization": f"Bearer {self.TOKEN}",
                "Content-Type": "application/json",
            },
        )
        response = connection.getresponse()
        result = json.loads(response.read())
        status = response.status
        connection.close()
        return status, result

    def test_render_and_extract_endpoints(self) -> None:
        status, rendered = self._post(
            "/api/v2/slides/render",
            {
                "presentation": dataclasses.asdict(
                    _presentation_with_inline_assets()
                ),
                "include_canonical_html": True,
            },
        )
        self.assertEqual(status, 200)
        self.assertTrue(rendered["success"])
        self.assertIn("current_ir", rendered)
        self.assertNotIn("<script", rendered["editable_html"].lower())

        status, extracted = self._post(
            "/api/v2/slides/extract-editable",
            {"canonical_html": rendered["canonical_html"]},
        )
        self.assertEqual(status, 200)
        self.assertTrue(extracted["canonical_ir_valid"])
        self.assertEqual(
            extracted["current_ir_sha256"],
            rendered["current_ir_sha256"],
        )

    def test_render_rejects_non_boolean_canonical_flag(self) -> None:
        status, payload = self._post(
            "/api/v2/slides/render",
            {
                "presentation": dataclasses.asdict(
                    _presentation_with_inline_assets()
                ),
                "include_canonical_html": "true",
            },
        )
        self.assertEqual(status, 400)
        self.assertIn("include_canonical_html", payload["error"])

    def test_browser_font_subset_endpoint_fails_closed_for_invalid_font(self) -> None:
        status, payload = self._post(
            "/api/v2/fonts/subset",
            {
                "font_data_base64": base64.b64encode(b"not-a-font").decode("ascii"),
                "characters": "中文 A",
                "face_index": 0,
            },
        )
        self.assertEqual(status, 422)
        self.assertFalse(payload["success"])
        self.assertEqual(
            payload["error_code"],
            "browser-font-unsupported-source",
        )

    def test_compile_endpoint_reconciles_real_asset_without_persisting_binary(self) -> None:
        image = b"\x89PNG\r\n\x1a\ncompile-composed"
        presentation = Presentation(
            width=960,
            height=540,
            slides=[
                Slide(
                    index=0,
                    shapes=[
                        Shape(
                            type="auto_shape",
                            name="Stable fill",
                            shape_id=17,
                            bounds=Bounds(12, 18, 320, 180),
                            fill=Fill(
                                type="picture",
                                image=ImageInfo(
                                    content_type="image/png",
                                    data=base64.b64encode(image).decode("ascii"),
                                    alpha=0.6,
                                ),
                            ),
                        )
                    ],
                )
            ],
        )
        status, rendered = self._post(
            "/api/v2/slides/render",
            {
                "presentation": dataclasses.asdict(presentation),
            },
        )
        self.assertEqual(status, 200)
        assets = {
            item["asset_id"]: {
                "content_type": item["content_type"],
                "data_base64": item["data_base64"],
            }
            for item in rendered["assets"]
        }
        edited = rendered["editable_html"].replace(
            'element-name="Stable fill"',
            'element-name="Edited fill"',
            1,
        )

        status, compiled = self._post(
            "/api/v2/slides/compile",
            {
                "editable_html": edited,
                "previous_current_ir": rendered["current_ir"],
                "presentation_assets": assets,
                "known_asset_ids": list(assets),
            },
        )

        self.assertEqual(status, 200, compiled)
        self.assertTrue(compiled["success"])
        self.assertEqual(compiled["reconcile_mode"], "stable-id-merge")
        serialized = json.dumps(compiled["current_ir"])
        self.assertIn("yrkie-asset-base64://", serialized)
        self.assertNotIn(";base64,", serialized)
        self.assertNotIn("data:image/", compiled["editable_html"])
        source_image = compiled["current_ir"]["slides"][0]["shapes"][0]["fill"]["image"]
        self.assertEqual(source_image["content_type"], "image/png")
        self.assertEqual(source_image["alpha"], 0.6)
        outer_svg_ids = {
            item["asset_id"] for item in rendered["assets"]
            if item["content_type"] == "image/svg+xml"
        }
        self.assertTrue(outer_svg_ids)
        self.assertTrue(
            any(asset_id in compiled["editable_html"] for asset_id in outer_svg_ids)
        )

    def test_render_endpoint_uses_transient_asset_for_composed_svg(self) -> None:
        image = b"\x89PNG\r\n\x1a\ntransient-composed"
        asset_id = hashlib.sha256(image).hexdigest()
        presentation = Presentation(
            width=960,
            height=540,
            slides=[
                Slide(
                    index=0,
                    shapes=[
                        Shape(
                            type="auto_shape",
                            shape_id=7,
                            bounds=Bounds(0, 0, 320, 180),
                            fill=Fill(
                                type="picture",
                                image=ImageInfo(
                                    content_type="image/png",
                                    data=f"yrkie-asset-base64://{asset_id}",
                                    alpha=0.5,
                                ),
                            ),
                        )
                    ],
                )
            ],
        )

        status, rendered = self._post(
            "/api/v2/slides/render",
            {
                "presentation": dataclasses.asdict(presentation),
                "known_asset_ids": [asset_id],
                "presentation_assets": {
                    asset_id: {
                        "content_type": "image/png",
                        "data_base64": base64.b64encode(image).decode("ascii"),
                    }
                },
            },
        )

        self.assertEqual(status, 200, rendered)
        self.assertNotIn("data:image/", rendered["editable_html"])
        self.assertNotIn(asset_id, {item["asset_id"] for item in rendered["assets"]})
        svg_asset = next(
            item for item in rendered["assets"]
            if item["content_type"] == "image/svg+xml"
        )
        svg = base64.b64decode(svg_asset["data_base64"]).decode("utf-8")
        self.assertIn("data:image/png;base64,", svg)

    def test_compile_persists_replaced_image_fit_crop_and_source_metadata(self) -> None:
        old_image = _png((220, 30, 30), (8, 4))
        new_image = _png((30, 80, 220), (8, 4))
        presentation = Presentation(
            width=960,
            height=540,
            slides=[
                Slide(
                    index=0,
                    shapes=[
                        Shape(
                            type="picture",
                            shape_id=31,
                            source_scope="slide",
                            bounds=Bounds(10, 10, 200, 200),
                            image=ImageInfo(
                                content_type="image/png",
                                data=base64.b64encode(old_image).decode("ascii"),
                                embed_id="rId31",
                                crop_left=0.1,
                            ),
                        )
                    ],
                )
            ],
        )
        status, rendered = self._post(
            "/api/v2/slides/render",
            {"presentation": dataclasses.asdict(presentation)},
        )
        self.assertEqual(status, 200, rendered)
        current = rendered["current_ir"]
        current["slides"][0]["shapes"][0]["image"]["future_image_metadata"] = {
            "source": "preserved"
        }
        old_asset_id = hashlib.sha256(old_image).hexdigest()
        new_asset_id = hashlib.sha256(new_image).hexdigest()
        edited = rendered["editable_html"].replace(
            f"yrkie-asset-data://{old_asset_id}",
            f"yrkie-asset-data://{new_asset_id}",
            1,
        )
        img_start = edited.index("<img ")
        img_end = edited.index(">", img_start)
        img_tag = edited[img_start:img_end]
        img_tag = img_tag.replace(
            ' style="',
            (
                ' data-crop-left="10" data-crop-top="20" '
                'data-crop-right="5" data-crop-bottom="15" style="'
            ),
            1,
        ).replace('style="', 'style="object-fit:cover;', 1)
        edited = edited[:img_start] + img_tag + edited[img_end:]
        assets = {
            item["asset_id"]: {
                "content_type": item["content_type"],
                "data_base64": item["data_base64"],
            }
            for item in rendered["assets"]
        }
        assets[new_asset_id] = {
            "content_type": "image/png",
            "data_base64": base64.b64encode(new_image).decode("ascii"),
        }

        status, compiled = self._post(
            "/api/v2/slides/compile",
            {
                "editable_html": edited,
                "previous_current_ir": current,
                "presentation_assets": assets,
                "known_asset_ids": list(assets),
            },
        )

        self.assertEqual(status, 200, compiled)
        self.assertEqual(compiled["reconcile_mode"], "stable-id-merge")
        image = compiled["current_ir"]["slides"][0]["shapes"][0]["image"]
        self.assertEqual(
            image["data"],
            f"yrkie-asset-base64://{new_asset_id}",
        )
        self.assertEqual(image["embed_id"], "rId31")
        self.assertEqual(
            image["future_image_metadata"],
            {"source": "preserved"},
        )
        self.assertIsNone(image["intrinsic_width_pt"])
        self.assertAlmostEqual(image["crop_left"], 0.10, places=4)
        self.assertAlmostEqual(image["crop_top"], 0.20, places=4)
        self.assertAlmostEqual(image["crop_right"], 0.05, places=4)
        self.assertAlmostEqual(image["crop_bottom"], 0.15, places=4)

    def test_compile_persists_cover_fit_without_explicit_crop(self) -> None:
        image = _png((20, 160, 80), (8, 4))
        presentation = Presentation(
            width=960,
            height=540,
            slides=[
                Slide(
                    index=0,
                    shapes=[
                        Shape(
                            type="picture",
                            shape_id=32,
                            bounds=Bounds(0, 0, 200, 200),
                            image=ImageInfo(
                                content_type="image/png",
                                data=base64.b64encode(image).decode("ascii"),
                            ),
                        )
                    ],
                )
            ],
        )
        status, rendered = self._post(
            "/api/v2/slides/render",
            {"presentation": dataclasses.asdict(presentation)},
        )
        self.assertEqual(status, 200, rendered)
        edited = rendered["editable_html"]
        img_start = edited.index("<img ")
        img_end = edited.index(">", img_start)
        img_tag = edited[img_start:img_end].replace(
            'style="',
            'style="object-fit:cover;',
            1,
        )
        edited = edited[:img_start] + img_tag + edited[img_end:]
        assets = {
            item["asset_id"]: {
                "content_type": item["content_type"],
                "data_base64": item["data_base64"],
            }
            for item in rendered["assets"]
        }

        status, compiled = self._post(
            "/api/v2/slides/compile",
            {
                "editable_html": edited,
                "previous_current_ir": rendered["current_ir"],
                "presentation_assets": assets,
                "known_asset_ids": list(assets),
            },
        )

        self.assertEqual(status, 200, compiled)
        image_ir = compiled["current_ir"]["slides"][0]["shapes"][0]["image"]
        self.assertGreater(image_ir["crop_left"], 0.2)
        self.assertGreater(image_ir["crop_right"], 0.2)
        self.assertAlmostEqual(image_ir["crop_top"], 0.0, places=4)
        self.assertAlmostEqual(image_ir["crop_bottom"], 0.0, places=4)

    def test_compile_persists_merged_table_cell_text_and_style(self) -> None:
        table = TableInfo(
            rows=[
                [
                    TableCell(
                        text="Merged",
                        colspan=2,
                        text_frame=_text_frame("Merged"),
                    ),
                    TableCell(text="", h_merge=True),
                ],
                [
                    TableCell(text="Cell C", text_frame=_text_frame("Cell C")),
                    TableCell(
                        text="Cell D",
                        text_frame=_text_frame("Cell D"),
                        border_left=Border(width=1, color="111111"),
                    ),
                ],
            ],
            col_widths=[150, 150],
            row_heights=[50, 50],
            style_id="source-style",
        )
        presentation = Presentation(
            width=960,
            height=540,
            slides=[
                Slide(
                    index=0,
                    shapes=[
                        Shape(
                            type="graphic_frame",
                            shape_id=44,
                            source_scope="slide",
                            bounds=Bounds(100, 100, 300, 100),
                            table=table,
                        )
                    ],
                )
            ],
        )
        status, rendered = self._post(
            "/api/v2/slides/render",
            {"presentation": dataclasses.asdict(presentation)},
        )
        self.assertEqual(status, 200, rendered)
        current = rendered["current_ir"]
        current["slides"][0]["shapes"][0]["table"]["future_table_metadata"] = {
            "source": "kept"
        }
        current["slides"][0]["shapes"][0]["table"]["rows"][1][1][
            "future_cell_metadata"
        ] = {"source": "kept"}
        edited = rendered["editable_html"].replace("Cell D", "Updated D", 1)
        text_position = edited.index("Updated D")
        cell_start = edited.rfind("<td", 0, text_position)
        cell_end = edited.index(">", cell_start)
        cell_tag = edited[cell_start:cell_end]
        if 'style="' in cell_tag:
            cell_tag = re.sub(
                r'style="([^"]*)"',
                (
                    r'style="\1;background-color:#336699;padding:8px 12px;'
                    r'border-left:2px dashed #ff0000"'
                ),
                cell_tag,
                count=1,
            )
        else:
            cell_tag += (
                ' style="background-color:#336699;padding:8px 12px;'
                'border:2px dashed #ff0000"'
            )
        edited = edited[:cell_start] + cell_tag + edited[cell_end:]

        status, compiled = self._post(
            "/api/v2/slides/compile",
            {
                "editable_html": edited,
                "previous_current_ir": current,
            },
        )

        self.assertEqual(status, 200, compiled)
        self.assertEqual(compiled["reconcile_mode"], "stable-id-merge")
        result_table = compiled["current_ir"]["slides"][0]["shapes"][0]["table"]
        self.assertEqual(result_table["style_id"], "source-style")
        self.assertEqual(
            result_table["future_table_metadata"],
            {"source": "kept"},
        )
        self.assertEqual(result_table["rows"][0][0]["colspan"], 2)
        self.assertTrue(result_table["rows"][0][1]["h_merge"])
        cell = result_table["rows"][1][1]
        self.assertEqual(cell["text"], "Updated D")
        self.assertEqual(
            cell["text_frame"]["paragraphs"][0]["runs"][0]["text"],
            "Updated D",
        )
        self.assertEqual(cell["fill"]["color"], "336699")
        self.assertAlmostEqual(cell["margins"]["left"], 9.0, places=2)
        self.assertEqual(cell["border_left"]["dash"], "dash")
        self.assertEqual(cell["border_left"]["color"], "FF0000")
        self.assertEqual(
            cell["future_cell_metadata"],
            {"source": "kept"},
        )


if __name__ == "__main__":
    unittest.main()
